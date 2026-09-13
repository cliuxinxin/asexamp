import asyncio
import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tcg.main import create_app


class Model:
    def __init__(self, questions=None, fail_task=None, delay_task=None, bad_refs=False):
        self.calls = []
        self.questions = questions or []
        self.fail_task = fail_task
        self.delay_task = delay_task
        self.bad_refs = bad_refs

    async def generate(self, task, context):
        self.calls.append((task, context))
        if task == self.delay_task:
            await asyncio.sleep(30)
        if task == self.fail_task:
            raise RuntimeError('Deliberate model failure')
        refs = [e['id'] for e in context.get('evidence', []) if e['role'] != 'example'][:1]
        if self.bad_refs:
            refs = ['invented-evidence']
        if task == 'route':
            return {'intent': 'generate_case'}
        if task == 'analyze_requirement':
            return {'items': [{'id': 'REQ-1', 'title': 'Authentication', 'description': 'Valid and invalid credentials', 'refs': refs}], 'report': {'questions': self.questions, 'assumptions': [], 'requirement_map': {'rules': ['valid credentials']}}}
        if task == 'generate_scenarios':
            return {'items': [{'id': 'SC-1', 'title': 'Login', 'description': 'Check authentication', 'priority': 'P1', 'refs': refs}], 'has_more': False}
        if task in ('generate_cases', 'import_cases'):
            scenarios = context.get('scenarios', [])
            return {'items': [{'id': 'TC-1', 'title': 'Login correctly', 'scenario_id': scenarios[0]['id'] if scenarios else '', 'type': 'Business', 'priority': 'P1', 'preconditions': 'Account exists', 'steps': [{'action': 'Enter valid credentials', 'expected': 'User is authenticated'}], 'refs': refs}], 'has_more': False}
        if task == 'review_cases':
            item = dict(context['cases'][0])
            item['title'] = 'Reviewed login case'
            return {'operations': [{'op': 'update', 'id': item['id'], 'item': item}], 'report': {'summary': 'One review completed'}}
        if task == 'query':
            return {'answer': 'Valid credentials authenticate the user.', 'refs': refs}
        if task == 'learn_template':
            return {'config': {'case_types': ['Business'], 'additional_rules': 'Use concise titles'}, 'summary': 'Learned sample format'}
        if task == 'modify':
            item = dict(context['artifact']['items'][0])
            item['title'] = 'Modified through chat'
            return {'operations': [{'op': 'update', 'id': item['id'], 'item': item}], 'summary': 'Updated title'}
        raise AssertionError(task)


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path, Model())) as client:
        yield client


def setup_chat(client, text='Users must log in with valid credentials. Invalid passwords must show an error without revealing whether the account exists.'):
    project = client.get('/api/projects').json()[0]
    chat = client.post(f"/api/projects/{project['id']}/chats", json={'title': 'Login'}).json()
    source = client.post(f"/api/chats/{chat['id']}/sources/text", json={'name': 'Requirements', 'text': text, 'role': 'primary'}).json()
    return project, chat, source


def start(client, chat, **kwargs):
    body = {'content': 'Generate cases', 'intent': 'generate_case', 'mode': 'auto', **kwargs}
    response = client.post(f"/api/chats/{chat['id']}/messages", json=body)
    assert response.status_code == 200, response.text
    return response.json()['run']


def until(client, run, statuses=('completed', 'failed', 'waiting', 'cancelled')):
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline:
        value = client.get(f"/api/runs/{run['id']}").json()
        if value['status'] in statuses:
            return value
        time.sleep(0.02)
    raise AssertionError(value)


def test_auto_publishes_only_reviewed_cases_and_exactly_one_review(tmp_path):
    model = Model()
    with TestClient(create_app(tmp_path, model)) as client:
        _, chat, source = setup_chat(client)
        run = until(client, start(client, chat))
        assert run['status'] == 'completed', run
        assert len(run['artifact_ids']) == 1
        artifact = client.get('/api/artifacts/' + run['artifact_ids'][0]).json()
        assert artifact['type'] == 'cases'
        assert artifact['items'][0]['title'] == 'Reviewed login case'
        assert sum(task == 'review_cases' for task, _ in model.calls) == 1
        messages = client.get('/api/chats/' + chat['id']).json()['messages']
        assert len(messages) == 2
        assert messages[-1]['metadata']['artifact_ids'] == run['artifact_ids']
        evidence = client.get('/api/sources/' + source['id']).json()
        assert artifact['items'][0]['refs'] == [evidence['chunks'][0]['id']]


def test_revision_conflict_restore_and_excel_formula_escaping(client):
    _, chat, _ = setup_chat(client)
    run = until(client, start(client, chat))
    aid = run['artifact_ids'][0]
    artifact = client.get('/api/artifacts/' + aid).json()
    items = artifact['items']
    items[0]['title'] = '=HYPERLINK("https://malicious.invalid")'
    updated = client.put('/api/artifacts/' + aid, json={'expected_revision': artifact['revision'], 'items': items})
    assert updated.status_code == 200
    assert client.put('/api/artifacts/' + aid, json={'expected_revision': artifact['revision'], 'items': items}).status_code == 409
    workbook = load_workbook(io.BytesIO(client.get('/api/artifacts/' + aid + '/export').content))
    assert workbook.active['B2'].value.startswith("'=")
    assert workbook.active['B2'].data_type == 's'
    restored = client.post('/api/artifacts/' + aid + '/restore', json={'revision': artifact['revision'], 'expected_revision': updated.json()['revision']})
    assert restored.status_code == 200
    assert restored.json()['revision'] > updated.json()['revision']
    assert restored.json()['items'][0]['title'] == 'Reviewed login case'
    assert client.get('/api/artifacts/' + aid + '/export?ids=unknown').status_code == 400


def test_project_boundaries_origin_and_invalid_refs(client):
    project, chat, _ = setup_chat(client)
    second = client.post('/api/projects', json={'name': 'Isolated'}).json()
    profile = client.get('/api/projects/' + second['id'] + '/profiles').json()[0]
    assert client.post('/api/chats/' + chat['id'] + '/messages', json={'content': 'Generate cases', 'intent': 'generate_case', 'mode': 'auto', 'profile_id': profile['id']}).status_code == 400
    assert client.post('/api/projects', json={'name': 'CSRF'}, headers={'Origin': 'https://evil.invalid'}).status_code == 403
    run = until(client, start(client, chat))
    artifact = client.get('/api/artifacts/' + run['artifact_ids'][0]).json()
    artifact['items'][0]['refs'] = ['not-real']
    assert client.put('/api/artifacts/' + artifact['id'], json={'expected_revision': artifact['revision'], 'items': artifact['items']}).status_code == 400
    artifact['items'][0]['refs'] = []
    assert client.put('/api/artifacts/' + artifact['id'], json={'expected_revision': artifact['revision'], 'items': artifact['items']}).status_code == 400


def test_document_failures_and_deleted_source_preserves_evidence(client):
    _, chat, source = setup_chat(client)
    endpoint = '/api/chats/' + chat['id'] + '/sources'
    assert client.post(endpoint, files={'file': ('payload.exe', b'executable')}).status_code == 400
    from pypdf import PdfWriter
    data = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(data)
    result = client.post(endpoint, files={'file': ('scan.pdf', data.getvalue(), 'application/pdf')})
    assert result.status_code == 400
    assert 'OCR' in result.json()['detail']
    assert client.delete('/api/sources/' + source['id']).status_code == 200
    assert client.get('/api/sources/' + source['id']).json()['chunks']
    assert client.get('/api/chats/' + chat['id']).json()['sources'] == []


def test_missing_model_and_missing_inputs_are_actionable(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        _, chat, _ = setup_chat(client)
        response = client.post('/api/chats/' + chat['id'] + '/messages', json={'content': 'Generate cases'})
        assert response.status_code == 400
        assert not client.get('/api/chats/' + chat['id']).json()['messages']
    with TestClient(create_app(tmp_path / 'injected', Model())) as client:
        project = client.get('/api/projects').json()[0]
        chat = client.post('/api/projects/' + project['id'] + '/chats', json={}).json()
        run = until(client, start(client, chat))
        assert run['status'] == 'completed'
        result = client.get('/api/artifacts/' + run['artifact_ids'][0]).json()
        assert result['type'] == 'answer'
        assert '需求' in result['items'][0]['description']


def test_second_server_cannot_claim_same_data_directory(tmp_path):
    with TestClient(create_app(tmp_path, Model())):
        with pytest.raises(Exception, match='目录|directory'):
            with TestClient(create_app(tmp_path, Model())):
                pass


@pytest.mark.parametrize(('intent', 'kind'), [('review_requirement', 'analysis'), ('generate_scenario', 'scenarios')])
def test_explicit_intent_stops_at_requested_final_artifact(client, intent, kind):
    _, chat, _ = setup_chat(client)
    run = until(client, start(client, chat, intent=intent))
    assert run['status'] == 'completed'
    assert len(run['artifact_ids']) == 1
    artifact = client.get('/api/artifacts/' + run['artifact_ids'][0]).json()
    assert artifact['type'] == kind


def test_native_office_upload_preserves_text_and_coordinates(client, tmp_path):
    from docx import Document
    from openpyxl import Workbook
    _, chat, _ = setup_chat(client)
    endpoint = '/api/chats/' + chat['id'] + '/sources'
    document = Document()
    document.add_paragraph('An account must have a unique email address.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Rule'
    table.cell(0, 1).text = 'Reject duplicate email'
    word = io.BytesIO()
    document.save(word)
    uploaded = client.post(endpoint, files={'file': ('requirements.docx', word.getvalue())})
    assert uploaded.status_code == 200, uploaded.text
    source = client.get('/api/sources/' + uploaded.json()['id']).json()
    assert 'unique email' in source['text'] and 'duplicate email' in source['text']
    assert any('表' in c['location'] for c in source['chunks'])
    workbook = Workbook()
    workbook.active.title = 'Rules'
    workbook.active.append(['Limit', 5])
    excel = io.BytesIO()
    workbook.save(excel)
    uploaded = client.post(endpoint, files={'file': ('requirements.xlsx', excel.getvalue())})
    assert uploaded.status_code == 200, uploaded.text
    source = client.get('/api/sources/' + uploaded.json()['id']).json()
    assert 'A1=Limit' in source['text'] and 'B1=5' in source['text']
    assert source['chunks'][0]['location'].startswith('Rules')


def test_arbitrary_case_fields_roundtrip_and_do_not_leak_into_export(client):
    _, chat, _ = setup_chat(client)
    run = until(client, start(client, chat))
    aid = run['artifact_ids'][0]
    artifact = client.get('/api/artifacts/' + aid).json()
    artifact['items'][0]['automation'] = {'owner': 'QA', 'ready': False}
    result = client.put('/api/artifacts/' + aid, json={'expected_revision': artifact['revision'], 'items': artifact['items']})
    assert result.status_code == 200
    assert result.json()['items'][0]['automation'] == {'owner': 'QA', 'ready': False}
    sheet = load_workbook(io.BytesIO(client.get('/api/artifacts/' + aid + '/export?layout=step').content)).active
    assert sheet.max_column == 7
    assert sheet.max_row == 2
