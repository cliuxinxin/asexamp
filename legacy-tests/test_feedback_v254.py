import asyncio, copy, io, json
import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tcg.main import create_app
from tcg.model import LangChainGateway, SYSTEM, TASK_INSTRUCTIONS
from tcg.schemas import DomainError
from test_gateway_v22 import configured, response, run_with_transport
from test_backend_api import setup_chat, start, until
from test_workflow_v25 import FlowModel


def test_complete_values_missing_only_final_brace_do_not_recall_model(tmp_path):
    value={'items':[{'id':'REQ-1','title':'登录'}],'report':{'diagrams':[{'mermaid':'mindmap\n  root((登录))'}]}}
    raw=json.dumps(value,ensure_ascii=False)[:-1]
    assert run_with_transport(configured(tmp_path),lambda r:httpx.Response(200,json=response(raw)))==value
    # Unfinished strings/values must still fail rather than inventing content.
    for bad in ('{"a":"unfinished','{"a":','{"a":12'):
        with pytest.raises(DomainError):
            run_with_transport(configured(tmp_path),lambda r:httpx.Response(200,json=response(bad)))


class ScopeModel(FlowModel):
    async def generate(self,task,context):
        if task=='generate_scenarios':
            self.calls.append((task,copy.deepcopy(context)))
            refs=[context['evidence'][0]['id']]
            return {'items':[{'id':sid,'title':title,'description':title,'priority':'P1','refs':refs,'requirement_ids':[context['analysis'][0]['id']]} for sid,title in [('SC-LOGIN','登录'),('SC-SCOPE','AI 引擎内部实现不在测试范围')]],'has_more':False}
        if task=='generate_cases':
            self.calls.append((task,copy.deepcopy(context)))
            return {'items':[{'id':'TC-'+s['id'],'title':s['title'],'scenario_id':s['id'],'type':'Business','priority':'P1','preconditions':'已有账号','test_data':'账号=pilot','steps':[{'action':'登录','expected':'显示入口'}],'refs':s['refs']} for s in context['scenarios']],'has_more':False}
        if task=='review_cases':
            self.calls.append((task,copy.deepcopy(context)))
            target=next(c for c in context['cases'] if 'SCOPE' in c['id'])
            report={'summary':'删除错误生成的范围排除项，保留登录用例。'}
            if context.get('validation_repair'):
                report['scenario_exclusions']=[{'scenario_id':target['scenario_id'],'reason':'需求明确排除 AI 引擎内部实现','refs':target['refs']}]
            return {'operations':[{'op':'delete','id':target['id']}],'report':report}
        return await super().generate(task,context)


def test_auto_progress_template_and_review_repair_are_observable(tmp_path):
    model=ScopeModel()
    async def handler(request):
        body=json.loads(request.content)
        task=next(k for k,v in TASK_INSTRUCTIONS.items() if SYSTEM+'\nTASK CONTRACT:\n'+v==body['messages'][0]['content'][0]['text'])
        context=json.loads(body['messages'][1]['content'][0]['text'])
        return httpx.Response(200,json=response(json.dumps(await model.generate(task,context),ensure_ascii=False)))
    transport=httpx.AsyncClient(transport=httpx.MockTransport(handler))
    app=create_app(tmp_path,LangChainGateway(configured(tmp_path),http_client=transport))
    try:
        with TestClient(app) as c:
            _,chat,_=setup_chat(c,'白名单账号登录后显示 AI Roleplay 入口。AI 引擎内部实现不在测试范围。')
            template=c.post('/api/chats/'+chat['id']+'/sources/text',json={'name':'用例字段定义','role':'example','text':'测试数据列必须写为字段=值。模板示例业务：购买火箭，不是本次需求。'}).json()
            columns=[{'field':'title','header':'场景标题'},{'field':'test_data','header':'测试数据','definition':'写为字段=值'},{'field':'expected','header':'验收结果'}]
            run=until(c,start(c,chat,experience='reliable',profile_override={'sheet_name':'验收表','excel_columns':columns,'template_rules':'预期必须可观察'}))
            assert run['status']=='completed',run.get('error')
            messages=c.get('/api/chats/'+chat['id']).json()['messages']
            stages=[m for m in messages if m.get('metadata',{}).get('stage_artifact_id')]
            assert [m['metadata']['stage'] for m in stages]==['understand','scenarios','cases','review']
            for message in stages:
                meta=message['metadata']
                assert c.get('/api/artifacts/'+meta['stage_artifact_id']+'/revisions/'+str(meta['stage_revision'])).status_code==200
            analysis=next(ctx for task,ctx in model.calls if task=='analyze_requirement')
            assert 'excel_columns' not in analysis['profile']
            assert not analysis.get('format_references')
            generation=next(ctx for task,ctx in model.calls if task=='generate_cases')
            assert generation['profile']['excel_columns']==columns
            assert generation['format_references'][0]['source_id']==template['id']
            assert '测试数据列' in generation['format_references'][0]['text']
            assert all(e['role']!='example' for e in generation['evidence'])
            artifact=c.get('/api/artifacts/'+run['artifact_ids'][-1]).json()
            assert len(artifact['items'])==1
            assert artifact['report']['template_usage']['columns']==columns
            assert artifact['report']['review_reports'][0]['scenario_exclusions']
            wb=load_workbook(io.BytesIO(c.get('/api/artifacts/'+artifact['id']+'/export').content))
            assert wb.active.title=='验收表'
            assert [cell.value for cell in wb.active[1]]==['场景标题','测试数据','验收结果']
            assert wb.active.cell(2,2).value=='账号=pilot'
            records=app.state.store.db.execute('SELECT call_id,payload FROM model_requests WHERE run_id=? ORDER BY rowid',(run['id'],)).fetchall()
            reviews=[row['call_id'] for row in records if json.loads(row['payload'])['task']=='review_cases']
            assert len(reviews)==2
            log=c.get('/api/runs/'+run['id']+'/failed-step',params={'call_id':reviews[0]}).text
            assert 'lost_scenario_coverage' in log
            assert '字段或业务校验未通过' in log
            assert 'JSON 语法：通过' in log
            # Old versions omitted call_id on schema errors; time/node association must recover them.
            for row in app.state.store.db.execute('SELECT id,payload FROM diagnostics WHERE run_id=?',(run['id'],)).fetchall():
                event=json.loads(row['payload'])
                if event.get('event')=='batch.validation_failed':
                    event.pop('call_id',None);event.pop('call_key',None)
                    app.state.store.db.execute('UPDATE diagnostics SET payload=? WHERE id=?',(json.dumps(event),row['id']))
            legacy_log=c.get('/api/runs/'+run['id']+'/failed-step',params={'call_id':reviews[0]}).text
            assert '字段或业务校验未通过' in legacy_log and 'lost_scenario_coverage' in legacy_log
    finally:asyncio.run(transport.aclose())


def test_case_fields_in_understanding_are_repaired_before_scenarios(tmp_path):
    class WrongStage(FlowModel):
        async def generate(self,task,context):
            result=await super().generate(task,context)
            if task=='analyze_requirement' and not context.get('validation_repair'):
                result['items'][0].update(steps='Step 1: 登录',expected='Result 1: 进入首页')
            return result
    model=WrongStage()
    with TestClient(create_app(tmp_path,model)) as c:
        _,chat,_=setup_chat(c)
        run=until(c,start(c,chat,experience='reliable'))
        assert run['status']=='completed',run.get('error')
        assert len([1 for task,_ in model.calls if task=='analyze_requirement'])==2
        ctx=next(ctx for task,ctx in model.calls if task=='generate_scenarios')
        assert 'steps' not in ctx['analysis'][0]


def test_review_cannot_exclude_a_scenario_using_an_invented_reference(tmp_path):
    class Ungrounded(ScopeModel):
        async def generate(self,task,context):
            result=await super().generate(task,context)
            if task=='review_cases' and context.get('validation_repair'):
                result['report']['scenario_exclusions'][0]['refs']=['invented-source#P1']
            return result
    with TestClient(create_app(tmp_path,Ungrounded())) as c:
        _,chat,_=setup_chat(c)
        run=until(c,start(c,chat,experience='reliable'))
        assert run['status']=='failed'
        assert any(e['code']=='invalid_exclusion' for e in run['validation_errors'])
        stages=c.get('/api/chats/'+chat['id']).json()['messages']
        draft=next(m for m in stages if m.get('metadata',{}).get('stage')=='cases')
        assert len(c.get('/api/artifacts/'+draft['metadata']['stage_artifact_id']).json()['items'])==2
