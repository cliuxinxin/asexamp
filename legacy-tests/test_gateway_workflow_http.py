"""End-to-end workflow through a test-only HTTP gateway, never production fallback."""
import asyncio
import io
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tcg.main import create_app
from tcg.model import SYSTEM, TASK_INSTRUCTIONS
from test_backend_api import start, until
from test_reliable_workflow import ReliableModel


def test_docx_to_xlsx_through_real_chat_completions_http(tmp_path):
    seen = []
    fixture = ReliableModel()

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            seen.append((self.path, self.headers.get('X-API-Key'), body))
            task = next(k for k, v in TASK_INSTRUCTIONS.items()
                        if SYSTEM + '\nTASK CONTRACT:\n' + v == body['messages'][0]['content'][0]['text'])
            context = json.loads(body['messages'][1]['content'][0]['text'])
            result = asyncio.run(fixture.generate(task, context))
            data = json.dumps({'choices': [{'message': {'content': json.dumps(result)}, 'finish_reason': 'stop'}]}).encode()
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Content-Length', str(len(data)))
            self.end_headers()
            self.wfile.write(data)

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        (tmp_path / '.env').write_text(
            f'TCG_MODEL_PROVIDER=openai\nTCG_MODEL_BASE_URL=http://127.0.0.1:{server.server_port}/api/v1\n'
            'TCG_MODEL_NAME=gpt-5\nTCG_API_KEY=TEST-ONLY-NOT-A-REAL-KEY\n')
        document = Document()
        document.add_paragraph('登录成功后进入首页。')
        document.add_paragraph('错误密码必须显示错误提示。')
        buffer = io.BytesIO()
        document.save(buffer)
        with TestClient(create_app(tmp_path)) as client:
            project = client.get('/api/projects').json()[0]
            chat = client.post('/api/projects/' + project['id'] + '/chats', json={'title': 'HTTP protocol test'}).json()
            upload = client.post('/api/chats/' + chat['id'] + '/sources', data={'role': 'auto'},
                                 files={'file': ('登录需求.docx', buffer.getvalue())})
            assert upload.status_code == 200, upload.text
            run = until(client, start(client, chat, experience='reliable', case_types=['Business']))
            assert run['status'] == 'completed', run
            exported = client.get('/api/artifacts/' + run['artifact_ids'][0] + '/export')
            assert exported.status_code == 200
            workbook = load_workbook(io.BytesIO(exported.content))
            assert workbook.active.max_row >= 2
        assert len(seen) >= 4
        for path, key, body in seen:
            assert path == '/api/v1/chat/completions'
            assert key == 'TEST-ONLY-NOT-A-REAL-KEY'
            assert set(body) == {'model', 'messages'}
            assert all(isinstance(m['content'], list) for m in body['messages'])
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
