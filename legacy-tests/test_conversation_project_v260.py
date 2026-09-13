"""Conversation project capabilities use real SQLite, template policies and Excel bytes."""
import asyncio
import copy
import io

import tempfile
import unittest
from contextlib import contextmanager
from openpyxl import load_workbook

from tcg.conversation_project import execute, register_routes, frozen_export
from tcg.documents import parse_text
from tcg.schemas import DomainError, MessageInput
from tcg.storage import Store


class ControlledEngine:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def fits(self, task, context):
        return True

    async def invoke_model(self, task, context, run_id=None):
        self.calls.append((task, copy.deepcopy(context)))
        return copy.deepcopy(self.response)


@contextmanager
def fixture():
    directory = tempfile.TemporaryDirectory()
    store = Store(directory.name)
    project = store.list('project')[0]
    chat = store.create_chat(project['id'], 'conversation')
    profile = store.list('profile')[0]
    try:
        yield store, chat, profile
    finally:
        store.close()
        directory.cleanup()


def source(store, chat, name='requirements', role='primary', text='失败五次锁定十分钟'):
    content, chunks = parse_text(text)
    return store.add_source(chat['id'], name, role, content, chunks)


def call(setup, name, args, engine=None):
    store, chat, _ = setup
    return asyncio.run(execute(store, engine or ControlledEngine({}), chat, name, args))


def artifacts(setup):
    store, chat, _ = setup
    src = source(store, chat)
    run = store.create_run(chat['id'], MessageInput(content='生成场景', intent='generate_scenario',
        mode='hitp', experience='reliable').model_dump())[1]
    ref = src['id'] + '#P1'
    analysis = store.artifact(run['id'], 'v7:analysis:artifact', 'analysis', '需求理解', [
        {'id': 'R1', 'title': '锁定', 'description': '十分钟', 'refs': [ref]},
        {'id': 'R2', 'title': '无关需求', 'description': '其他业务', 'refs': [ref]},
    ])
    scenarios = store.artifact(run['id'], 'v4:scenarios_artifact', 'scenarios', '场景', [
        {'id': 'S1', 'title': '锁定', 'priority': 'P1', 'description': '失败五次', 'requirement_ids': ['R1'], 'refs': [ref]},
        {'id': 'S2', 'title': '无关场景', 'priority': 'P2', 'description': '其他业务', 'requirement_ids': ['R2'], 'refs': [ref]},
    ], {'lineage': {'analysis_artifact_id': analysis['id'], 'analysis_revision': 1}})
    cases = store.artifact(run['id'], 'cases', 'cases', '用例', [
        {'id': 'C1', 'title': '锁定', 'scenario_id': 'S1', 'priority': 'P1', 'type': 'Negative',
         'preconditions': '已注册', 'steps': [{'action': '连续失败五次', 'expected': '锁定十分钟'}],
         'human_notes': '人工备注', 'refs': [ref]},
    ], {'lineage': {'scenario_artifact_id': scenarios['id'], 'scenario_revision': 1}})
    for artifact in (analysis, scenarios, cases):
        store.put('artifact', {**artifact, '_visible': True})
    store.update_run(run['id'], status='waiting', interrupt={'type': 'scenario_review',
        'artifact_id': scenarios['id']}, _interrupt_id='confirm-scenarios', stop_after='scenarios')
    return run, analysis, scenarios, cases


def dual_template():
    return {'template_kinds': ['scenarios', 'cases'], 'summary': '独立提取两种模板', 'config': {
        'scenario_excel_columns': [{'field': 'title', 'header': '场景名称'}, {'field': 'description', 'header': '场景说明'}],
        'scenario_sheet_name': '已学场景', 'scenario_filename_pattern': 'scenarios_{project}.xlsx',
        'excel_columns': [{'field': 'title', 'header': '用例名称', 'value_source': 'ai', 'required': True},
            {'field': 'steps', 'header': '操作', 'value_source': 'derived'},
            {'field': 'expected', 'header': '预期', 'value_source': 'derived'},
            {'field': 'human_notes', 'header': '人工备注', 'value_source': 'manual'}],
        'sheet_name': '已学用例', 'filename_pattern': 'cases_{project}.xlsx'}}


def test_learn_apply_both_without_generation_run_and_preserve_manual_policy(setup):
    store, chat, profile = setup
    example = source(store, chat, '双模板', 'example', '场景名称、场景说明；用例名称、操作、预期、人工备注')
    engine = ControlledEngine(dual_template())
    result = call(setup, 'project.learn_template', {'source_ids': [example['id']], 'kind': 'both', 'apply': True}, engine)
    assert result['status'] == 'succeeded'
    saved = store.get('profile', profile['id'])
    assert saved['version'] == 2
    assert saved['config']['scenario_sheet_name'] == '已学场景'
    assert saved['config']['sheet_name'] == '已学用例'
    assert saved['config']['excel_columns'][-1]['value_source'] == 'manual'
    assert len(store.runs(chat_id=chat['id'])) == 0
    assert engine.calls[0][1]['format_references'][0]['id'] == example['id'] + '#P1'
    assert store.get('chat', chat['id'])['profile_id'] == profile['id']


def test_learning_only_persists_proposal_and_apply_uses_current_profile_version(setup):
    store, chat, profile = setup
    example = source(store, chat, '模板', 'example')
    learned = call(setup, 'project.learn_template', {'source_ids': [example['id']]}, ControlledEngine(dual_template()))
    assert store.get('profile', profile['id'])['version'] == 1
    template_id = learned['templates'][0]['id']
    updated = store.update_profile(profile['id'], profile['name'], {**profile['config'], 'scope': '更新的业务范围'}, 1)
    with unittest.TestCase().assertRaisesRegex(DomainError, 'Profile'):
        call(setup, 'project.apply_profile', {'template_ids': [template_id], 'expected_version': 1})
    applied = call(setup, 'project.apply_profile', {'template_ids': [template_id], 'expected_version': updated['version']})
    assert applied['profile']['version'] == 3
    assert store.get('profile', profile['id'])['config']['scope'] == '更新的业务范围'


def test_source_roles_project_isolation_and_existing_snapshot_not_rewritten(setup):
    store, chat, _ = setup
    run, *_ = artifacts(setup)
    before = copy.deepcopy(store.run(run['id']))
    added = call(setup, 'project.add_sources', {'content': '管理员可强制下线', 'role': 'change'})
    created = store.get('source', added['sources'][0]['id'])
    assert created['role'] == 'change' and created['project_id'] == chat['project_id']
    assert store.run(run['id'])['_source_ids'] == before['_source_ids']
    other = store.create_project('other')
    other_chat = store.create_chat(other['id'], 'other')
    outside = source(store, other_chat)
    with unittest.TestCase().assertRaisesRegex(DomainError, '项目'):
        call(setup, 'project.add_sources', {'source_ids': [outside['id']]})
    assert store.get('source', outside['id'])['role'] == 'primary'


def test_exports_two_actual_frozen_files_with_independent_templates(setup):
    store, chat, profile = setup
    _, _, scenarios, cases = artifacts(setup)
    store.update_profile(profile['id'], profile['name'], {**profile['config'], **dual_template()['config']}, 1)
    result = call(setup, 'artifact.export', {'artifact_ids': [scenarios['id'], cases['id']], 'profile_id': profile['id']})
    files = next(p['files'] for p in result['parts'] if p['type'] == 'files')
    assert len(files) == 2 and files[0]['url'] != files[1]['url']
    changed = copy.deepcopy(scenarios['items']); changed[0]['title'] = '导出后的新标题'
    store.revise_artifact(scenarios['id'], 1, changed)
    books = [load_workbook(io.BytesIO(frozen_export(store, f['url'].rsplit('/', 1)[-1])[1])) for f in files]
    assert books[0].sheetnames == ['已学场景']
    assert list(books[0].active.values)[0] == ('场景名称', '场景说明')
    assert list(books[0].active.values)[1][0] == '锁定'
    assert books[1].sheetnames == ['已学用例']
    assert list(books[1].active.values)[0] == ('用例名称', '操作', '预期', '人工备注')
    assert list(books[1].active.values)[1][-1] == '人工备注'


def test_pin_rejects_stale_revision_and_new_chat_generation_uses_shared_samples(setup):
    store, chat, profile = setup
    _, _, _, cases = artifacts(setup)
    with unittest.TestCase().assertRaisesRegex(DomainError, '版本'):
        call(setup, 'project.pin_samples', {'artifact_id': cases['id'], 'expected_revision': 2, 'selected_ids': ['C1']})
    result = call(setup, 'project.pin_samples', {'artifact_id': cases['id'], 'expected_revision': 1, 'selected_ids': ['C1']})
    assert result['profile']['version'] == 2
    second = store.create_chat(chat['project_id'], 'second')
    src = source(store, second)
    run = store.create_run(second['id'], MessageInput(content='新业务', intent='generate_case', profile_id=profile['id']).model_dump())[1]
    assert run['_profile']['sample_cases'][0]['steps'] == cases['items'][0]['steps']
    assert 'scenario_id' not in run['_profile']['sample_cases'][0]
    assert 'refs' not in run['_profile']['sample_cases'][0]
    assert run['_source_ids'] == [src['id']]


def test_template_cannot_reclassify_manual_execution_fields_as_ai_and_waiting_run_uses_new_version(setup):
    store, chat, profile = setup
    run, *_ = artifacts(setup)
    profile = store.update_profile(profile['id'], profile['name'], {**profile['config'], 'excel_columns': [
        {'field': 'human_notes', 'header': '人工备注', 'value_source': 'manual', 'required': False}]}, 1)
    example = source(store, chat, '模板', 'example')
    response = dual_template()
    response['config']['excel_columns'] += [
        {'field': 'actual_result', 'header': '实际结果', 'value_source': 'ai', 'required': True}]
    response['config']['excel_columns'][-2].update(value_source='ai', required=True)
    result = call((store, chat, profile), 'project.learn_template', {'source_ids': [example['id']], 'apply': True}, ControlledEngine(response))
    saved = store.get('profile', profile['id'])
    assert result['profile']['version'] == 3
    columns = {c['field']: c for c in saved['config']['excel_columns']}
    assert columns['human_notes']['value_source'] == 'manual'
    assert columns['actual_result']['value_source'] == 'manual'
    current = store.run(run['id'])
    assert current['_profile'] == saved['config']
    assert current['status'] == 'waiting' and current['stop_after'] == 'scenarios'


def test_explicit_snapshot_export_preserves_artifact_template_after_chat_profile_change(setup):
    store, chat, profile = setup
    _, _, scenarios, _ = artifacts(setup)
    store.update_profile(profile['id'], profile['name'], {**profile['config'], **dual_template()['config']}, 1)
    store.put('chat', {**chat, 'profile_id': profile['id']})
    result = call(setup, 'artifact.export', {'artifact_id': scenarios['id'], 'revision': 1, 'use_snapshot': True})
    file = result['parts'][0]['files'][0]
    book = load_workbook(io.BytesIO(frozen_export(store, file['url'].rsplit('/', 1)[-1])[1]))
    assert book.sheetnames == ['Test Scenarios']
    assert list(book.active.values)[0][0] == 'Scenario ID'


def test_success_receipt_prevents_duplicate_profile_application(setup):
    store, chat, profile = setup
    example = source(store, chat, '模板', 'example')
    store.put('conversation_command', {'id': 'command-test', 'project_id': chat['project_id'],
        'chat_id': chat['id'], 'status': 'running'})
    engine = ControlledEngine(dual_template())
    args = {'source_ids': [example['id']], 'apply': True}
    first = asyncio.run(execute(store, engine, chat, 'project.learn_template', args, 'command-test'))
    second = asyncio.run(execute(store, engine, chat, 'project.learn_template', args, 'command-test'))
    assert first == second
    assert store.get('profile', profile['id'])['version'] == 2
    assert len(engine.calls) == 1
    assert store.get('conversation_command', 'command-test')['result'] == first


class SourceImpactEngine(ControlledEngine):
    async def invoke_model(self, task, context, run_id=None):
        self.calls.append((task, copy.deepcopy(context)))
        if task == 'project_source_impact':
            return {'requirement_ids': ['R1'], 'summary': '锁定规则增加管理员强制下线',
                    'global_impact': False, 'uncertain': False, 'refs': [context['new_evidence'][0]['id']]}
        if task == 'artifact_modify':
            assert [r['id'] for r in context['artifact']['items']] == ['R1']
            ref = next(e['id'] for e in context['evidence'] if e['role'] == 'change')
            return {'operations': [{'op': 'update', 'id': 'R1', 'item': {'description': '锁定十分钟；管理员可强制下线', 'refs': [ref]}}], 'summary': '更新锁定需求'}
        if task == 'artifact_sync_scenarios':
            assert context['analysis'][0]['description'] == '锁定十分钟；管理员可强制下线'
            return {'operations': [{'op': 'update', 'id': 'S1', 'item': {'description': '验证锁定与管理员强制下线'}}], 'summary': '定向同步场景'}
        raise AssertionError('不得启动无关生成：' + task)


def test_targeted_sources_updates_analysis_and_only_related_scenarios_preserving_run(setup):
    store, chat, _ = setup
    run, analysis, scenarios, cases = artifacts(setup)
    src = source(store, chat, '强制下线变更', 'change', '管理员可强制下线')
    engine = SourceImpactEngine({})
    result = call(setup, 'project.update_from_sources', {'source_ids': [src['id']],
        'artifact_id': analysis['id'], 'targets': ['analysis', 'scenarios']}, engine)
    assert result['status'] == 'succeeded'
    updated_analysis = store.get('artifact', analysis['id'])
    updated_scenarios = store.get('artifact', scenarios['id'])
    assert updated_analysis['revision'] == updated_scenarios['revision'] == 2
    assert updated_analysis['items'][1] == analysis['items'][1]
    assert updated_scenarios['items'][1] == scenarios['items'][1]
    assert store.get('artifact', cases['id'])['revision'] == 1
    assert store.revision(analysis['id'], 1)['items'] == analysis['items']
    current = store.run(run['id'])
    assert current['status'] == 'waiting' and current['stop_after'] == 'scenarios'
    assert src['id'] in current['_source_ids']
    assert len(store.runs(chat_id=chat['id'])) == 1
    assert [task for task, _ in engine.calls] == ['project_source_impact', 'artifact_modify', 'artifact_sync_scenarios']


def test_global_source_impact_requests_scope_and_does_not_change_saved_artifacts(setup):
    store, chat, _ = setup
    _, analysis, _, _ = artifacts(setup)
    src = source(store, chat, '全局规则', 'change', '全部会话都执行新认证规则')
    engine = ControlledEngine({'requirement_ids': [], 'summary': '无法限定在单一需求，需扩大全局分析',
        'global_impact': True, 'uncertain': True, 'refs': [src['id'] + '#P1']})
    result = call(setup, 'project.update_from_sources', {'source_ids': [src['id']], 'artifact_id': analysis['id']}, engine)
    assert result['status'] == 'needs_input'
    assert store.get('artifact', analysis['id'])['revision'] == 1
    assert len(engine.calls) == 1


class ConversationProjectTests(unittest.TestCase):
    pass


def _test(function):
    def run(self):
        with fixture() as setup:
            function(setup)
    return run


for _name, _function in list(globals().items()):
    if _name.startswith('test_'):
        setattr(ConversationProjectTests, _name, _test(_function))
        del globals()[_name]


if __name__ == '__main__':
    unittest.main()
