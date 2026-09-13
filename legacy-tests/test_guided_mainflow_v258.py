"""One HTTP main flow: uploaded template, grounded guidance, HITP and both XLSX exports."""
import copy
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from tcg.main import create_app
from test_backend_api import start, until
from test_workflow_v25 import FlowModel


QUESTION = '允许哪些角色登录？'
ANSWER = '已注册用户可以使用有效凭证登录。'
CASE_COLUMNS = [
    {'field': 'title', 'header': '用例名称'},
    {'field': 'expected', 'header': '预期结果'},
]
SCENARIO_COLUMNS = [
    {'field': 'title', 'header': '场景名称'},
    {'field': 'description', 'header': '场景说明'},
    {'field': 'priority', 'header': '优先级'},
    {'field': 'requirement_ids', 'header': '需求编号'},
    {'field': 'refs', 'header': '证据引用'},
]


class GuidedMainFlowModel(FlowModel):
    def __init__(self):
        super().__init__(questions=[QUESTION])
        self.example_ref = None

    async def generate(self, task, context):
        if task == 'learn_template':
            self.calls.append((task, copy.deepcopy(context)))
            return {
                'template_kinds': ['scenarios'],
                'config': {
                    'scenario_excel_columns': SCENARIO_COLUMNS,
                    'scenario_sheet_name': '业务场景',
                    'scenario_filename_pattern': '{project}_learned_scenarios_{date}.xlsx',
                },
                'summary': '识别场景模板的工作表和列顺序；示例业务内容不属于需求。',
            }
        result = await super().generate(task, context)
        if task == 'analyze_requirement':
            ref = next(item['id'] for item in context['evidence'] if item['role'] == 'primary')
            suggestion = {
                'question': QUESTION, 'answer': ANSWER,
                'basis': '需求明确说明已注册用户使用有效凭证登录。',
                'refs': [ref], 'confidence': 'supported',
            }
            result['report']['question_suggestions'] = [
                suggestion,
                {**suggestion, 'answer': '模板中的退款审批角色可以登录。', 'refs': [self.example_ref]},
            ]
        return result


def test_uploaded_scenario_template_guided_hitp_and_reviewed_exports(tmp_path):
    model = GuidedMainFlowModel()
    with TestClient(create_app(tmp_path, model)) as client:
        project = client.get('/api/projects').json()[0]
        created = client.post('/api/projects/' + project['id'] + '/chats', json={'title': '登录主流程'})
        assert created.status_code == 200, created.text
        chat = created.json()
        endpoint = '/api/chats/' + chat['id'] + '/sources'
        requirement = client.post(
            endpoint, data={'role': 'primary'},
            files={'file': ('login-requirements.txt',
                            (ANSWER + '无效密码应显示错误，且不得泄露账号是否存在。').encode('utf-8'),
                            'text/plain')},
        )
        assert requirement.status_code == 200, requirement.text
        requirement_id = requirement.json()['id']
        primary_refs = {item['id'] for item in client.get('/api/sources/' + requirement_id).json()['chunks']}

        workbook = Workbook()
        workbook.active.title = '业务场景'
        workbook.active.append([column['header'] for column in SCENARIO_COLUMNS])
        workbook.active.append(['退款审批示例', '仅展示填写方式，不作为本轮需求', 'P1', 'EXAMPLE-REQ', 'example-only'])
        template_bytes = BytesIO()
        workbook.save(template_bytes)
        template = client.post(
            endpoint, data={'role': 'example'},
            files={'file': ('scenario-template.xlsx', template_bytes.getvalue(),
                            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        )
        assert template.status_code == 200, template.text
        template_id = template.json()['id']
        template_source = client.get('/api/sources/' + template_id).json()
        assert template_source['role'] == 'example'
        assert '场景名称' in template_source['text'] and '退款审批示例' in template_source['text']
        model.example_ref = template_source['chunks'][0]['id']

        profile_response = client.post('/api/projects/' + project['id'] + '/profiles', json={
            'name': '场景和用例独立模板',
            'config': {'excel_columns': CASE_COLUMNS, 'sheet_name': '既有用例',
                       'filename_pattern': '{project}_existing_cases_{date}.xlsx'},
        })
        assert profile_response.status_code == 200, profile_response.text
        profile = profile_response.json()
        before_profiles = client.get('/api/projects/' + project['id'] + '/profiles').json()
        learned = until(client, start(
            client, chat, experience='reliable', intent='learn_template',
            content='请学习已上传的 Excel 场景模板。', source_ids=[template_id], profile_id=profile['id'],
        ))
        assert learned['status'] == 'completed', learned
        proposal = client.get('/api/artifacts/' + learned['artifact_ids'][0]).json()
        assert proposal['type'] == 'proposal'
        assert proposal['report']['template_kinds'] == ['scenarios']
        proposed_config = proposal['report']['config']
        assert proposed_config['scenario_excel_columns'] == SCENARIO_COLUMNS
        assert proposed_config['excel_columns'] == CASE_COLUMNS
        learned_messages = client.get('/api/chats/' + chat['id']).json()['messages']
        learned_message = next(message for message in learned_messages
                               if message['role'] == 'assistant'
                               and message['metadata']['run_id'] == learned['id'])
        assert learned_message['metadata']['proposal'] == {
            'config': proposed_config, 'template_kinds': ['scenarios'],
        }
        assert client.get('/api/projects/' + project['id'] + '/profiles').json() == before_profiles
        learning_context = next(context for task, context in model.calls if task == 'learn_template')
        assert any(item['id'] == model.example_ref and '场景名称' in item['text']
                   for item in learning_context['evidence'])

        # The existing proposal action explicitly saves report.config through the profile API.
        saved = client.put('/api/profiles/' + profile['id'], json={
            'name': profile['name'], 'config': proposed_config, 'expected_version': profile['version'],
        })
        assert saved.status_code == 200, saved.text
        saved_config = saved.json()['config']
        assert saved_config['scenario_excel_columns'] == SCENARIO_COLUMNS
        assert saved_config['scenario_sheet_name'] == '业务场景'
        assert saved_config['excel_columns'] == profile['config']['excel_columns']
        assert saved_config['sheet_name'] == profile['config']['sheet_name']
        assert saved_config['filename_pattern'] == profile['config']['filename_pattern']
        assert 'template_kinds' not in saved_config

        run = until(client, start(client, chat, experience='reliable', mode='hitp', profile_id=profile['id']))
        assert run['status'] == 'waiting' and run['interrupt']['type'] == 'clarification', run
        suggestions = run['interrupt']['question_suggestions']
        assert len(suggestions) == 1, suggestions
        suggestion = suggestions[0]
        assert suggestion['question'] == QUESTION and suggestion['answer'] == ANSWER
        assert suggestion['confidence'] == 'supported' and suggestion['basis']
        assert suggestion['refs'] and set(suggestion['refs']) <= primary_refs
        assert model.example_ref not in suggestion['refs']
        assert not any(task == 'generate_cases' for task, _ in model.calls)

        resumed = client.post('/api/runs/' + run['id'] + '/resume', json={'answer': suggestion['answer']})
        assert resumed.status_code == 200, resumed.text
        run = until(client, run)
        assert run['status'] == 'waiting' and run['interrupt']['type'] == 'strategy_review', run
        analysis = client.get('/api/artifacts/' + run['interrupt']['artifact_id']).json()
        assert analysis['report']['questions'] == []
        assert ANSWER in analysis['report']['clarification']
        resumed = client.post('/api/runs/' + run['id'] + '/resume', json={'approved': True})
        assert resumed.status_code == 200, resumed.text
        run = until(client, run)
        assert run['status'] == 'waiting' and run['interrupt']['type'] == 'scenario_review', run
        scenario = client.get('/api/artifacts/' + run['interrupt']['artifact_id']).json()
        assert scenario['type'] == 'scenarios'
        assert not any(task == 'generate_cases' for task, _ in model.calls)
        resumed = client.post('/api/runs/' + run['id'] + '/resume', json={'approved': True})
        assert resumed.status_code == 200, resumed.text
        run = until(client, run)
        assert run['status'] == 'completed', run
        cases = client.get('/api/artifacts/' + run['artifact_ids'][-1]).json()
        assert cases['type'] == 'cases' and cases['items'][0]['title'] == 'Reviewed login case'
        assert sum(task == 'analyze_requirement' for task, _ in model.calls) == 1
        assert sum(task == 'review_cases' for task, _ in model.calls) == 1
        scenario_context = next(context for task, context in model.calls if task == 'generate_scenarios')
        confirmed_map = scenario_context['global_requirement_map']
        assert ANSWER in confirmed_map['clarification']
        assert confirmed_map['previous_questions'] == [QUESTION]
        sources = client.get('/api/chats/' + chat['id']).json()['sources']
        assert sum(source['role'] == 'clarification' for source in sources) == 1

        calls_before_export = len(model.calls)
        scenario_export = client.get('/api/artifacts/' + scenario['id'] + '/export')
        assert scenario_export.status_code == 200, scenario_export.text
        scenario_sheet = load_workbook(BytesIO(scenario_export.content)).active
        assert scenario_sheet.title == '业务场景'
        assert list(next(scenario_sheet.values)) == [column['header'] for column in SCENARIO_COLUMNS]
        assert scenario_sheet.max_row == len(scenario['items']) + 1
        assert scenario_sheet.cell(2, 1).value == scenario['items'][0]['title']
        assert scenario_sheet.cell(2, 4).value == '\n'.join(scenario['items'][0]['requirement_ids'])
        assert scenario_sheet.cell(2, 5).value == '\n'.join(scenario['items'][0]['refs'])
        assert 'learned_scenarios' in scenario_export.headers['content-disposition']
        case_export = client.get('/api/artifacts/' + cases['id'] + '/export')
        assert case_export.status_code == 200, case_export.text
        case_sheet = load_workbook(BytesIO(case_export.content)).active
        assert case_sheet.title == '既有用例'
        assert list(next(case_sheet.values)) == ['用例名称', '预期结果']
        assert case_sheet.cell(2, 1).value == 'Reviewed login case'
        assert case_sheet.cell(2, 2).value == '1. User is authenticated'
        assert 'existing_cases' in case_export.headers['content-disposition']
        assert len(model.calls) == calls_before_export
