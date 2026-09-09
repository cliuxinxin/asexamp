"""Local text extraction with source coordinates and safe XLSX export."""
import csv
import io
import re
import zipfile
from pathlib import Path

from .schemas import DomainError

MAX_UPLOAD = 100 * 1024 * 1024
MAX_TEXT = 2_000_000
ALLOWED = {'.txt', '.md', '.csv', '.docx', '.pdf', '.xlsx', '.xls'}


def split_parts(parts):
    chunks = []
    for location, text in parts:
        text = text.strip()
        if not text:
            continue
        # Keep paragraph and row boundaries; split exceptionally long paragraphs.
        for position in range(0, len(text), 2500):
            chunks.append({'text': text[position:position + 2500], 'location': location + (f' · 字符 {position + 1}' if position else '')})
    if not chunks:
        raise DomainError('文件未提取到可用文本；扫描 PDF 需要先使用本地 OCR 工具转换为文本 PDF')
    text = '\n\n'.join(chunk['text'] for chunk in chunks)
    if len(text) > MAX_TEXT:
        raise DomainError('解析文本超过 200 万字符；请拆分文档后上传，未截断或保存内容')
    return text, chunks


def parse_text(text):
    if not isinstance(text, str) or not text.strip():
        raise DomainError('来源文本不能为空')
    return split_parts([(f'段落 {i}', value) for i, value in enumerate(re.split(r'\n\s*\n', text), 1)])


def classify_source(name, text):
    """Conservative, explainable provisional role; user correction remains authoritative."""
    title = Path(name).stem.lower()
    for role, words in [('example', ('样例', '示例', 'sample', 'example')),
                        ('change', ('变更', '修订说明', 'change', 'release note')),
                        ('clarification', ('澄清', 'clarification')),
                        ('supplement', ('补充', 'supplement')),
                        ('knowledge', ('知识', '操作手册', 'knowledge', 'handbook'))]:
        if any(word in title for word in words):
            return role, {'label': role, 'reason': '根据文件标题识别用途，请确认是否正确', 'provisional': True}
    return 'primary', {'label': 'primary', 'reason': '暂归主要需求；无法仅凭内容可靠确定权威性，请按实际用途确认', 'provisional': True}


def parse_with_docling(name, data):
    """Optional local parser; weights must be installed before an offline demonstration."""
    import tempfile
    try:
        from docling.document_converter import DocumentConverter
    except ImportError:
        raise DomainError('尚未安装Docling。请安装requirements-docling.txt并准备模型，或使用原生解析支持的文本文件') from None
    with tempfile.TemporaryDirectory(prefix='tcg-docling-') as directory:
        path = Path(directory) / Path(name).name
        path.write_bytes(data)
        result = DocumentConverter().convert(path)
        status = getattr(result.status, 'value', str(result.status))
        if status != 'success':
            raise DomainError('Docling未完整解析该文件；未接受部分解析结果，请检查扫描质量或转换文件')
        parts = []
        for index, (item, _) in enumerate(result.document.iterate_items(), 1):
            if hasattr(item, 'export_to_markdown'):
                text = item.export_to_markdown(doc=result.document)
            else:
                text = getattr(item, 'text', '')
            if not text:
                continue
            pages = sorted({p.page_no for p in getattr(item, 'prov', [])})
            parts.append((f'Docling 区块 {index}' + (f' · 页 {pages}' if pages else ''), text))
        return split_parts(parts)


def parse_document(name, data, max_upload=MAX_UPLOAD, parser='native'):
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED:
        raise DomainError('不支持此格式；请上传 DOCX、文本 PDF、XLSX、XLS、CSV、TXT 或 MD')
    if len(data) > max_upload:
        raise DomainError(f'文件超过 {max_upload // (1024 * 1024)} MB 限制', 413)
    if suffix in ('.docx', '.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise DomainError('压缩文档展开后过大；请拆分文件')
        except zipfile.BadZipFile:
            raise DomainError('文件不是有效的 Office 文档') from None
    try:
        if parser == 'docling' and suffix in ('.pdf', '.docx', '.xlsx'):
            return parse_with_docling(name, data)
        if suffix in ('.txt', '.md', '.csv'):
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                try:
                    text = data.decode('gb18030')
                except UnicodeDecodeError:
                    raise DomainError('无法识别文本编码，请转为 UTF-8 后重试') from None
            if '\x00' in text:
                raise DomainError('文件包含二进制数据，请检查文本编码')
            if suffix == '.csv':
                return split_parts([(f'CSV 行 {i}', ' | '.join(row)) for i, row in enumerate(csv.reader(io.StringIO(text)), 1)])
            return parse_text(text)
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise DomainError('不支持加密 PDF，请先解密后上传')
            parts = []
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ''
                if not text.strip():
                    raise DomainError(f'PDF 第 {index} 页没有可提取文本；本地版本不内置 OCR，请先转换扫描页后上传')
                parts.append((f'PDF 第 {index} 页', text))
            return split_parts(parts)
        if suffix == '.docx':
            from docx import Document
            document = Document(io.BytesIO(data))
            parts = []
            for index, block in enumerate(document.iter_inner_content(), 1):
                if hasattr(block, 'text'):
                    parts.append((f'DOCX 段落 {index}', block.text))
                else:
                    parts.extend((f'DOCX 表 {index} 行 {row_number}', ' | '.join(cell.text for cell in row.cells)) for row_number, row in enumerate(block.rows, 1))
            return split_parts(parts)
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                parts = []
                for sheet in workbook:
                    for row in sheet.iter_rows():
                        values = [f'{cell.coordinate}={cell.value}' for cell in row if cell.value is not None]
                        if values:
                            parts.append((f'{sheet.title} · 行 {row[0].row}', ' | '.join(values)))
                return split_parts(parts)
            finally:
                workbook.close()
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data)
        return split_parts([(f'{sheet.name} · 行 {index + 1}', ' | '.join(str(value) for value in sheet.row_values(index))) for sheet in workbook.sheets() for index in range(sheet.nrows)])
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError(f'无法解析文件（{type(exc).__name__}）；请检查文件完整性或转换为 TXT/CSV') from None


def cell_safe(value):
    text = str(value) if value is not None else ''
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
    if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')):
        text = "'" + text
    return text


def export_cases(artifact, layout='case', selected=None):
    if artifact['type'] != 'cases':
        raise DomainError('仅 Case Artifact 支持 Excel 导出')
    if layout not in ('case', 'step'):
        raise DomainError('layout 必须为 case 或 step')
    items = artifact['items']
    if selected is not None:
        if not selected or not set(selected).issubset({item['id'] for item in items}):
            raise DomainError('导出所选条目 ID 无效')
        items = [item for item in items if item['id'] in set(selected)]
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    title = artifact.get('_profile', {}).get('sheet_name', 'Test Cases')
    sheet.title = re.sub(r'[\\/*?:\[\]]', '_', title)[:31] or 'Test Cases'
    sheet.append(['Case ID', 'Title', 'Type', 'Priority', 'Preconditions', 'Steps', 'Expected Result'])
    for item in items:
        base = [item['id'], item['title'], item['type'], item['priority'], item['preconditions']]
        if layout == 'step':
            for index, step in enumerate(item['steps'], 1):
                sheet.append([cell_safe(value) for value in base + [f'{index}. {step["action"]}', step['expected']]])
        else:
            actions = '\n'.join(f'{index}. {step["action"]}' for index, step in enumerate(item['steps'], 1))
            expected = '\n'.join(f'{index}. {step["expected"]}' for index, step in enumerate(item['steps'], 1))
            sheet.append([cell_safe(value) for value in base + [actions, expected]])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='23635B')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for column, width in {'A': 20, 'B': 42, 'C': 18, 'D': 12, 'E': 38, 'F': 64, 'G': 64}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
