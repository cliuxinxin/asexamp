import io

import pytest
from openpyxl import load_workbook

from tcg.documents import export_artifact
from tcg.schemas import DEFAULT_PROFILE, DomainError, profile_config
from tcg.workflow import valid_question_suggestions


def test_question_suggestions_keep_only_current_grounded_answers():
    evidence = {
        'src#P1': {'id': 'src#P1', 'role': 'primary'},
        'example#P1': {'id': 'example#P1', 'role': 'example'},
    }
    supported = {
        'question': '是否锁定账号？', 'answer': '连续失败 5 次后锁定。',
        'basis': '需求明确给出阈值。', 'refs': ['src#P1'], 'confidence': 'supported',
    }
    report = {'questions': ['是否锁定账号？'], 'question_suggestions': [
        supported,
        {**supported, 'question': '其他问题？'},
        {**supported, 'refs': ['example#P1']},
        {**supported, 'answer': ''},
    ]}
    assert valid_question_suggestions(report, evidence) == [supported]
    assert valid_question_suggestions({'questions': ['需人工回答？']}, evidence) == []


def test_question_suggestions_ignore_malformed_optional_data_and_duplicates():
    evidence = {'src#P1': {'id': 'src#P1', 'role': 'primary'}}
    suggestion = {'question': '是否锁定？', 'answer': '请确认是否按五次锁定。',
                  'basis': '这是待确认的假设；登录策略中提到失败次数。', 'refs': ['src#P1'], 'confidence': 'assumption'}
    assert valid_question_suggestions({'questions': ['是否锁定？'], 'question_suggestions': None}, evidence) == []
    report = {'questions': ['是否锁定？'], 'question_suggestions': [
        {**suggestion, 'question': {'text': '是否锁定？'}}, suggestion, suggestion,
        {**suggestion, 'refs': ['missing#P1']}, {**suggestion, 'confidence': 'certain'},
    ]}
    assert valid_question_suggestions(report, evidence) == [suggestion]


def test_malformed_suggestions_do_not_block_legacy_clarification(tmp_path):
    from fastapi.testclient import TestClient
    from tcg.main import create_app
    from test_backend_api import setup_chat, start, until
    from test_workflow_v25 import FlowModel

    class MalformedSuggestionModel(FlowModel):
        async def generate(self, task, context):
            result = await super().generate(task, context)
            if task == 'analyze_requirement':
                result['report']['question_suggestions'] = None
            return result

    with TestClient(create_app(tmp_path, MalformedSuggestionModel(questions=['是否锁定？']))) as client:
        _, chat, _ = setup_chat(client)
        run = until(client, start(client, chat, experience='reliable', mode='hitp'))
        assert run['status'] == 'waiting', run
        assert run['interrupt']['questions'] == ['是否锁定？']
        suggestions = run['interrupt']['question_suggestions']
        assert len(suggestions) == 1 and suggestions[0]['question'] == '是否锁定？'
        assert suggestions[0]['confidence'] == 'assumption' and suggestions[0]['refs'] == []
        assert '待确认' in suggestions[0]['answer']


def test_scenario_profile_is_independent_from_case_profile():
    profile = profile_config({'scenario_excel_columns': [{'field': 'title', 'header': '场景名称'}]})
    assert profile['scenario_sheet_name'] == 'Test Scenarios'
    assert profile['scenario_filename_pattern'] == '{project}_scenarios_{date}.xlsx'
    assert profile['scenario_excel_columns'] == [{'field': 'title', 'header': '场景名称'}]
    assert profile['excel_columns'] == DEFAULT_PROFILE['excel_columns']
    profile_config({'scenario_excel_columns': [
        {'field': 'requirement_ids', 'header': '需求编号'},
        {'field': 'refs', 'header': '证据'},
    ]})
    with pytest.raises(DomainError, match='来源或内部记录'):
        profile_config({'excel_columns': [{'field': 'refs', 'header': '证据'}]})


def test_scenario_export_default_custom_selected_and_safe_cells():
    artifact = {
        'type': 'scenarios',
        'items': [
            {'id': 'S-1', 'title': '=危险标题', 'description': '正常', 'priority': 'P1',
             'requirement_ids': ['R-1', 'R-2'], 'refs': ['src#P1']},
            {'id': 'S-2', 'title': '第二场景', 'description': '边界', 'priority': 'P2',
             'requirement_ids': ['R-2'], 'refs': ['src#P2']},
        ],
        '_profile': {},
    }
    sheet = load_workbook(io.BytesIO(export_artifact(artifact))).active
    assert sheet.title == 'Test Scenarios'
    assert list(next(sheet.values)) == ['Scenario ID', 'Title', 'Description', 'Priority', 'Requirement IDs', 'Evidence Refs']
    assert sheet.cell(2, 2).value == "'=危险标题"
    assert sheet.cell(2, 5).value == 'R-1\nR-2'

    custom = {**artifact, '_profile': {
        'scenario_sheet_name': '业务场景',
        'scenario_excel_columns': [
            {'field': 'title', 'header': '场景名称'},
            {'field': 'requirement_ids', 'header': '需求追溯'},
        ],
    }}
    selected = load_workbook(io.BytesIO(export_artifact(custom, selected=['S-2']))).active
    assert selected.title == '业务场景'
    assert list(next(selected.values)) == ['场景名称', '需求追溯']
    assert selected.max_row == 2 and selected.cell(2, 1).value == '第二场景'


def test_case_export_contract_is_unchanged():
    artifact = {'type': 'cases', 'items': [{
        'id': 'C-1', 'title': '登录', 'type': 'Business', 'priority': 'P1',
        'scenario_id': 'S-1', 'preconditions': '存在账号',
        'steps': [{'action': '登录', 'expected': '进入首页'}], 'refs': ['src#P1'],
    }], '_profile': {}}
    sheet = load_workbook(io.BytesIO(export_artifact(artifact))).active
    assert sheet.title == 'Test Cases'
    assert list(next(sheet.values))[:4] == ['Case ID', 'Title', 'Type', 'Priority']
