"""Viewer exports read the selected saved revision without model calls or edits."""
import copy
import io
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tcg.main import create_app
from tcg.storage import dump, now


@pytest.fixture
def viewer(tmp_path):
    app = create_app(tmp_path, model_gateway=SimpleNamespace())
    with TestClient(app, base_url='http://localhost') as client:
        store = app.state.store
        project = store.list('project')[0]
        chat = store.create_chat(project['id'], 'Export viewer')
        profile = {'excel_columns': [
            {'field': 'id', 'header': '编号'}, {'field': 'title', 'header': '标题'},
            {'field': 'steps', 'header': '操作'}, {'field': 'expected', 'header': '预期'}],
            'sheet_name': 'Old Cases', 'excel_layout': 'step',
            'scenario_excel_columns': [{'field': 'id', 'header': '场景编号'},
                {'field': 'title', 'header': '场景标题'}], 'scenario_sheet_name': 'Old Scenarios'}
        rows = [{'id': cid, 'title': 'Old ' + cid, 'type': 'Business', 'priority': 'P1',
            'scenario_id': 'S1', 'preconditions': 'Account exists', 'refs': [],
            'steps': [{'action': 'Login', 'expected': 'Success'},
                      {'action': 'Open page', 'expected': 'Visible'}]} for cid in ('C1', 'C2')]
        for aid, kind, items in [('cases', 'cases', rows), ('scenarios', 'scenarios',
            [{'id': 'S1', 'title': 'Old S1', 'requirement_ids': ['R1']}])]:
            historical = {'id': aid, 'project_id': project['id'], 'chat_id': chat['id'],
                'type': kind, 'title': aid, 'revision': 1, 'items': items,
                '_visible': True, '_profile': profile, 'report': {}}
            latest = copy.deepcopy(historical)
            latest['revision'] = 2
            for row in latest['items']:
                row['title'] = row['title'].replace('Old ', 'New ')
            latest['_profile'].update(sheet_name='New Cases', excel_layout='case',
                                      scenario_sheet_name='New Scenarios')
            store.put('artifact', latest)
            for version in (historical, latest):
                store.db.execute('INSERT INTO revisions VALUES(?,?,?,?,?,?)',
                    (aid, version['revision'], dump(version), now(), 'test fixture', '{}'))
        yield client, store, project


def workbook(response):
    assert response.status_code == 200, response.text
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats')
    return load_workbook(io.BytesIO(response.content)).active


def test_case_export_uses_selected_historical_rows_and_template_layout(viewer):
    client, store, _ = viewer
    before = copy.deepcopy(store.get('artifact', 'cases'))
    options = client.get('/api/artifacts/cases/export-options?revision=1').json()
    assert options['snapshot']['sheet_name'] == 'Old Cases'
    assert options['revision'] == 1
    sheet = workbook(client.get('/api/artifacts/cases/export?revision=1&ids=C2'))
    assert sheet.title == 'Old Cases'
    assert list(sheet.values) == [('编号', '标题', '操作', '预期'),
        ('C2', 'Old C2', '1. Login', '1. Success'),
        ('C2', 'Old C2', '2. Open page', '2. Visible')]
    assert store.get('artifact', 'cases') == before


def test_current_export_all_selected_and_explicit_profile(viewer):
    client, store, project = viewer
    sheet = workbook(client.get('/api/artifacts/cases/export'))
    assert sheet.title == 'New Cases'
    assert sheet.max_row == 3
    assert sheet.cell(2, 2).value == 'New C1'
    chosen = store.create_profile(project['id'], 'Export columns', {
        'excel_columns': [{'field': 'title', 'header': '验证名称'}], 'sheet_name': 'Chosen'})
    sheet = workbook(client.get('/api/artifacts/cases/export', params={
        'revision': 1, 'ids': 'C1', 'layout': 'case', 'profile_id': chosen['id']}))
    assert sheet.title == 'Chosen'
    assert list(sheet.values) == [('验证名称',), ('Old C1',)]


def test_scenario_export_uses_saved_revision_and_rejects_missing_revision(viewer):
    client, _, _ = viewer
    options = client.get('/api/artifacts/scenarios/export-options?revision=1').json()
    assert options['snapshot']['scenario_sheet_name'] == 'Old Scenarios'
    sheet = workbook(client.get('/api/artifacts/scenarios/export?revision=1'))
    assert sheet.title == 'Old Scenarios'
    assert list(sheet.values) == [('场景编号', '场景标题'), ('S1', 'Old S1')]
    for suffix in ('export', 'export-options'):
        assert client.get('/api/artifacts/scenarios/' + suffix + '?revision=99').status_code == 404
    assert client.get('/api/artifacts/cases/export?revision=1&ids=unknown').status_code == 400
