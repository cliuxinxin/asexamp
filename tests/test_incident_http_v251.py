"""HTTP protocol + user JSON replay, including a deliberately malformed response."""
import asyncio,io,json,zipfile
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from threading import Thread
from docx import Document
from openpyxl import Workbook,load_workbook
from fastapi.testclient import TestClient
from tcg.main import create_app
from tcg.model import SYSTEM,TASK_INSTRUCTIONS
from test_backend_api import start,until
from test_incident_v251 import IncidentModel


def test_http_hitp_repair_export_template_and_debug_bundle(tmp_path):
    model=IncidentModel();seen=[];malformed=False
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,*args):pass
        def do_POST(self):
            nonlocal malformed
            body=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            task=next(k for k,v in TASK_INSTRUCTIONS.items() if SYSTEM+'\nTASK CONTRACT:\n'+v==body['messages'][0]['content'][0]['text'])
            context=json.loads(body['messages'][1]['content'][0]['text']);seen.append((task,context,body))
            if task=='generate_scenarios' and not malformed:
                content='{"items": [}';malformed=True
            else:content=json.dumps(asyncio.run(model.generate(task,context)),ensure_ascii=False)
            data=json.dumps({'choices':[{'message':{'content':content},'finish_reason':'stop'}]}).encode()
            self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler);thread=Thread(target=server.serve_forever,daemon=True);thread.start()
    secret='TEST-ONLY-NOT-A-REAL-KEY'
    try:
        (tmp_path/'.env').write_text(f'TCG_MODEL_PROVIDER=openai\nTCG_MODEL_BASE_URL=http://127.0.0.1:{server.server_port}/api/v1\nTCG_MODEL_NAME=test\nTCG_API_KEY={secret}\n')
        document=Document();document.add_paragraph('用户登录、错误提示、5次失败锁定15分钟、成功登录清零计数以及退出后的保护页面访问。')
        data=io.BytesIO();document.save(data)
        with TestClient(create_app(tmp_path)) as c:
            project=c.get('/api/projects').json()[0];chat=c.post('/api/projects/'+project['id']+'/chats',json={}).json()
            assert c.post('/api/chats/'+chat['id']+'/sources',data={'role':'primary'},files={'file':('需求.docx',data.getvalue())}).status_code==200
            run=until(c,start(c,chat,experience='reliable',mode='hitp'))
            assert run['interrupt']['type']=='clarification',run
            aid=run['interrupt']['artifact_id']
            c.post('/api/runs/'+run['id']+'/resume',json={'answer':'不用扩展，按现有规则验证'})
            run=until(c,run);assert run['interrupt']['type']=='strategy_review',run
            assert run['interrupt']['artifact_id']==aid
            c.post('/api/runs/'+run['id']+'/resume',json={'approved':True})
            run=until(c,run);assert run['interrupt']['type']=='scenario_review',run.get('error')
            scenarios=c.get('/api/artifacts/'+run['interrupt']['artifact_id']).json()
            assert len(scenarios['items'])==20
            c.post('/api/runs/'+run['id']+'/resume',json={'approved':True})
            run=until(c,run);assert run['status']=='completed',run.get('error')
            artifact=c.get('/api/artifacts/'+run['artifact_ids'][-1]).json();assert artifact['type']=='cases'
            exported=c.get('/api/artifacts/'+artifact['id']+'/export');assert exported.status_code==200
            assert load_workbook(io.BytesIO(exported.content)).active.max_row==21
            assert sum(t=='analyze_requirement' for t,_,_ in seen)==1
            assert any(t=='generate_scenarios' and ctx.get('json_repair') for t,ctx,_ in seen)
            bundle=c.get('/api/runs/'+run['id']+'/debug-bundle');assert bundle.status_code==200,bundle.text[:300] if bundle.status_code!=200 else ''
            with zipfile.ZipFile(io.BytesIO(bundle.content)) as z:
                all_contents='\n'.join(z.read(n).decode() for n in z.namelist())
                assert '{"items": [}' in all_contents
                assert secret not in all_contents
                metadata=json.loads(z.read('diagnostics.json'))
                assert 'checkpoint' in metadata and metadata['checkpoint']['next']==[]
                assert any(e['event']=='json.repair_started' for e in metadata['events'])
            # A new generation over the same sources reuses the saved understanding.
            again=until(c,start(c,chat,experience='reliable',mode='auto'))
            assert again['status']=='completed',again.get('error')
            assert sum(t=='analyze_requirement' for t,_,_ in seen)==1
            wb=Workbook();wb.active.append(['Team','Members','Tokens']);wb.active.append(['A',8,123]);xlsx=io.BytesIO();wb.save(xlsx)
            assert c.post('/api/chats/'+chat['id']+'/sources',data={'role':'example'},files={'file':('SBU_Report.xlsx',xlsx.getvalue())}).status_code==200
            before=c.get('/api/projects/'+project['id']+'/profiles').json()
            learned=until(c,start(c,chat,experience='reliable',intent='learn_template'))
            assert learned['status']=='completed',learned.get('error')
            assert c.get('/api/projects/'+project['id']+'/profiles').json()==before
    finally:
        server.shutdown();server.server_close();thread.join(timeout=2)
