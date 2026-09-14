"""Resolved table cells survive the real review gate and queued Excel export."""
import copy
import io
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tcg.main import create_app
from test_native_journey_v300 import NativeJourneyGateway
from test_supervisor_http_v312 import QueueGateway, post, wait_files
from test_supervisor_pipeline_v312 import await_gate, execution_plan, plan_state


class TableReviewGateway(NativeJourneyGateway):
    """All generated statements are supported by the synthetic login source."""
    def __init__(self, count=3):
        super().__init__()
        self.count = count

    async def generate_native(self, task, context, schema, instruction):
        if task == 'generate_cases':
            result = await super().generate_native(task, context, schema, instruction)
            seed = result['items'][0]
            result['items'] = [{**copy.deepcopy(seed), 'id': f'TC-{index:03}',
                'title': f'登录验证组合 {index:02}', 'test_data': f'注册账号 account_{index:02}',
                'tester': '', 'steps': [
                    {'action': '打开登录页面', 'expected': '页面显示账号和密码输入框'},
                    {'action': f'填写有效账号 account_{index:02} 和密码', 'expected': '输入框接受填写值'},
                    {'action': '点击登录', 'expected': '登录成功'}]}
                for index in range(1, self.count + 1)]
            return result
        if task == 'review_cases':
            self.generations.append((task, copy.deepcopy(context)))
            rows = copy.deepcopy(context['cases'])
            rows[0]['title'] = 'AI 建议：明确有效凭证登录验证'
            rows[0]['steps'][2]['expected'] = '登录成功并显示首页'
            if len(rows) > 1:
                rows[1]['preconditions'] = '有效注册账号存在，用户尚未登录。'
            return {'items': rows, 'report': {'summary': '建议明确可验证的登录结果，并补足执行前提。',
                'issues': [{'case_id': rows[0]['id'], 'field': 'expected',
                    'title': '登录成功应明确可观察的首页结果。', 'severity': 'warning'},
                    {'title': '本次仅使用演示需求，账号均为合成数据。', 'severity': 'info'}]}}
        return await super().generate_native(task, context, schema, instruction)


SYNTHETIC_SOURCE = ('注册用户输入有效账号密码后登录成功并显示首页。登录页显示账号和密码输入框，'
    '输入框接受填写值。有效注册账号已存在，用户尚未登录。测试账号为 account_01 至 account_12。')


def configure_table_profile(store, project_id):
    profile = store.list('profile', project_id=project_id)[0]
    config = {**profile['config'], 'excel_layout': 'case', 'excel_columns': [
        {'field': 'title', 'header': '测试用例名称'},
        {'field': 'id', 'header': '用例编号'},
        {'field': 'preconditions', 'header': '前置条件'},
        {'field': 'test_data', 'header': '输入数据', 'value_source': 'ai', 'required': False},
        {'field': 'steps', 'header': '操作步骤', 'value_source': 'derived'},
        {'field': 'expected', 'header': '预期结果', 'value_source': 'derived'},
        {'field': 'tester', 'header': '执行人', 'value_source': 'manual', 'required': False}]}
    return store.update_profile(profile['id'], '全屏评审演示模板', config, profile['version'])


def begin_review_plan(client, app, gateway):
    project = client.get('/api/projects').json()[0]
    profile = configure_table_profile(app.state.store, project['id'])
    response = client.post('/api/projects/' + project['id'] + '/chats', json={'title': '全屏表格评审验收'})
    assert response.status_code == 200, response.text
    chat = response.json()
    upload = client.post('/api/chats/' + chat['id'] + '/sources',
        files={'file': ('synthetic-login.md', SYNTHETIC_SOURCE.encode(), 'text/markdown')},
        data={'role': 'primary'})
    assert upload.status_code == 200, upload.text
    gateway.plans = [execution_plan(
        ('pipeline_start', '根据需求生成用例，逐步人工确认，并生成 AI 评审建议。'),
        ('export', '评审确认完成后，导出最终确认的全部用例。'))]
    gateway.calls = [('start_pipeline_tool', {'mode': 'hitp'}), ('export_artifact_tool', {})]
    initial, _ = post(client, chat, 'table-review-start', '生成测试用例，每一步让我确认，评审完成后导出 Excel。',
        source_ids=[upload.json()['id']], profile_id=profile['id'])
    assert initial['status'] == 'succeeded', initial
    part = next(part for part in initial['parts'] if part['type'] == 'execution_plan')
    for index, kind in enumerate(('strategy_review', 'scenario_review')):
        run, prompt = await_gate(client, chat['id'], kind)
        response, _ = post(client, chat, f'table-review-gate-{index}', '同意', reply_kind='confirm',
            reply_to=prompt['id'], command={'name': 'workflow.resume',
                'arguments': {'run_id': run['id'], 'action': 'approved'}})
        assert response['status'] == 'succeeded', response
    run, prompt = await_gate(client, chat['id'], 'case_result_review')
    assert not app.state.store.list('frozen_export', chat_id=chat['id'])
    return chat, profile, run, prompt, part


def test_table_resolution_uses_real_gate_and_exports_exact_resolved_cells_once(tmp_path):
    gateway = QueueGateway()
    gateway.business_model = TableReviewGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        chat, profile, run, prompt, plan = begin_review_plan(client, app, gateway)
        base = '/api/artifacts/' + prompt['artifact_id'] + '/table-review'
        response = client.get(base, params={'run_id': run['id'], 'proposal_id': prompt['proposal_id']})
        assert response.status_code == 200, response.text
        view = response.json()
        assert not view['read_only'] and not view['stale'], view
        assert view['prompt_id'] == prompt['id']
        assert view['profile_id'] == profile['id'] and view['profile_revision'] == profile['version']
        assert [column['header'] for column in view['columns']] == [
            column['header'] for column in profile['config']['excel_columns']]
        assert len(view['issues']) == 2, view
        resolved = copy.deepcopy(view['proposed_items'])
        # Reject the title change, retain the suggested expected result, and enter a manual premise.
        resolved[0]['title'] = view['original_items'][0]['title']
        resolved[0]['preconditions'] = '已登录测试系统，待验证账号已注册。'
        resolved[2]['tester'] = '人工评审员'
        body = {'items': resolved, 'expected_revision': view['artifact_revision'],
            'profile_id': view['profile_id'], 'profile_revision': view['profile_revision'],
            'run_id': run['id'], 'proposal_id': view['proposal_id'], 'prompt_id': view['prompt_id'],
            'layout': 'case'}
        projection = client.post(base + '/project', json=body)
        assert projection.status_code == 200, projection.text
        projected = projection.json()
        draft_export = client.post(base + '/export', json=body)
        assert draft_export.status_code == 200, draft_export.text
        draft_values = list(load_workbook(io.BytesIO(draft_export.content)).active.values)
        expected = [[column['header'] for column in projected['columns']]] + [
            row['cells'] for row in projected['rows']]
        # openpyxl represents empty-string cells as None when reading back.
        assert [[value if value is not None else '' for value in row] for row in draft_values] == expected
        assert app.state.store.run(run['id'])['status'] == 'waiting'
        assert app.state.store.get('artifact', view['artifact_id'])['revision'] == view['artifact_revision']
        assert not app.state.store.list('frozen_export', chat_id=chat['id'])
        save_body = {**body, 'client_request_id': 'human-table-review-final'}
        saved = client.post(base + '/save', json=save_body)
        assert saved.status_code == 200, saved.text
        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            completed = app.state.store.run(run['id'])
            assert completed['status'] != 'failed', completed
            if completed['status'] == 'completed':
                break
            time.sleep(.02)
        assert completed['status'] == 'completed', completed
        files = wait_files(client, chat, saved.json())
        assert len(files) == 1
        artifact = app.state.store.get('artifact', view['artifact_id'])
        assert artifact['revision'] == view['artifact_revision'] + 1
        for saved_item, resolved_item in zip(artifact['items'], resolved):
            assert {key: value for key, value in saved_item.items() if not key.startswith('_') and key != 'refs'} == {
                key: value for key, value in resolved_item.items() if not key.startswith('_') and key != 'refs'}
        assert artifact['items'][0]['steps'][2]['expected'] == '登录成功并显示首页'
        assert artifact['items'][0]['refs'] != view['original_items'][0]['refs']
        records = app.state.store.list('frozen_export', chat_id=chat['id'])
        assert len(records) == 1 and records[0]['revision'] == artifact['revision']
        exported_values = list(load_workbook(io.BytesIO(client.get(files[0]['url']).content)).active.values)
        assert exported_values == draft_values
        assert plan_state(client, chat['id'], plan['plan_id'])['status'] == 'completed'
        assert gateway.executed == ['start_pipeline_tool', 'export_artifact_tool']
        assert len(gateway.planner_inputs) == 1
        duplicate = client.post(base + '/save', json=save_body)
        assert duplicate.status_code == 200, duplicate.text
        assert app.state.store.get('artifact', view['artifact_id'])['revision'] == artifact['revision']
        assert len(app.state.store.list('frozen_export', chat_id=chat['id'])) == 1
