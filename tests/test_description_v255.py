import copy,io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tcg.main import create_app
from tcg.documents import export_cases
from tcg.schemas import DomainError
from test_workflow_v25 import FlowModel
from test_backend_api import setup_chat,start,until

COLUMNS=[{'field':'title','header':'用例标题'},{'field':'case_description','header':'用例描述','definition':'说明验证目标、条件及预期行为'}]

class DescriptionModel(FlowModel):
    async def generate(self,task,context):
        if task=='complete_case_fields':
            self.calls.append((task,copy.deepcopy(context)))
            return {'items':[{'id':row['id'],'fields':{'description':'验证已有账号使用正确凭证登录后，能够成功进入首页。'},'unresolved':{}} for row in context['cases']]}
        return await super().generate(task,context)

def test_template_description_is_completed_and_exported(tmp_path):
    model=DescriptionModel()
    with TestClient(create_app(tmp_path,model)) as c:
        _,chat,_=setup_chat(c)
        run=until(c,start(c,chat,experience='reliable',profile_override={'excel_columns':COLUMNS}))
        assert run['status']=='completed',run.get('error')
        artifact=c.get('/api/artifacts/'+run['artifact_ids'][-1]).json()
        assert artifact['items'][0].get('description','').strip()
        assert sum(task=='complete_case_fields' for task,_ in model.calls)==1
        repair=next(ctx for task,ctx in model.calls if task=='complete_case_fields')
        assert repair['columns'][0]['definition']==COLUMNS[1]['definition']
        assert all(e['role']!='example' for e in repair['evidence'])
        result=c.get('/api/artifacts/'+artifact['id']+'/export')
        assert result.status_code==200
        sheet=load_workbook(io.BytesIO(result.content)).active
        assert sheet.cell(1,2).value=='用例描述'
        assert sheet.cell(2,2).value==artifact['items'][0]['description']

def test_existing_cases_can_fill_only_missing_descriptions(tmp_path):
    model=DescriptionModel()
    with TestClient(create_app(tmp_path,model)) as c:
        _,chat,_=setup_chat(c)
        run=until(c,start(c,chat,experience='reliable'))
        original=c.get('/api/artifacts/'+run['artifact_ids'][-1]).json()
        # Two existing cases, one already has a user-authored description.
        rows=copy.deepcopy(original['items']);rows.append({**copy.deepcopy(rows[0]),'id':'keep-me','description':'用户原有描述，不能改写。'})
        before=c.put('/api/artifacts/'+original['id'],json={'expected_revision':original['revision'],'items':rows}).json()
        marker=len(model.calls)
        requested=c.post('/api/artifacts/'+original['id']+'/complete-descriptions',json={})
        assert requested.status_code==200,requested.text
        completed=until(c,requested.json()['run'])
        assert completed['status']=='completed',completed.get('error')
        after=c.get('/api/artifacts/'+original['id']).json()
        assert after['items'][0]['description'].strip()
        assert after['items'][1]==before['items'][1]
        assert {k:v for k,v in after['items'][0].items() if k!='description'}==before['items'][0]
        assert [task for task,_ in model.calls[marker:]]==['complete_case_fields']
        assert c.post('/api/artifacts/'+original['id']+'/complete-descriptions',json={}).json()['unchanged'] is True

def test_description_export_uses_existing_alias_but_never_duplicates_title():
    case={'id':'C1','title':'标题','case_description':'原有描述','steps':[{'action':'登录','expected':'首页'}]}
    artifact={'type':'cases','items':[case],'_profile':{'excel_columns':[{'field':'description','header':'用例描述'}]}}
    sheet=load_workbook(io.BytesIO(export_cases(artifact))).active
    assert sheet.cell(2,1).value=='原有描述'
    case.pop('case_description')
    with pytest.raises(DomainError,match='用例描述'):export_cases(artifact)
