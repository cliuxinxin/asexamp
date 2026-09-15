"""Selected-scope continuation and approval crash-window regression coverage."""
import asyncio
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tcg.main import create_app
from tcg.table_review import TableSave, save_review
from test_supervisor_http_v312 import QueueGateway, begin, file_parts, plan, post, seed, wait_files
from workspace_helpers import save_workspace


def test_workspace_save_keeps_original_tail_selection_after_chat_selection_changes(tmp_path):
    gateway = QueueGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        chat, cases = asyncio.run(seed(app.state.store, gateway))
        initial, part, _ = begin(client, app, gateway, chat, cases)
        prompt = initial['pending'][0]
        guidance, _ = post(client, chat, 'natural-accept', '同意',
            reply_to=prompt['id'], artifact_id=cases['id'], artifact_revision=cases['revision'],
            selected_ids=[cases['items'][1]['id']])
        assert guidance['status'] == 'needs_confirmation'
        assert not app.state.store.list('frozen_export', chat_id=chat['id'])
        result, _ = save_workspace(client, prompt)
        files = wait_files(client, chat, result)
        sheet = load_workbook(io.BytesIO(client.get(files[0]['url']).content)).active
        values = list(sheet.values)
        assert len(values) == 2
        assert cases['items'][0]['id'] in values[1], values
        assert cases['items'][1]['id'] not in values[1]
        assert len(app.state.store.list('frozen_export', chat_id=chat['id'])) == 1
        assert len(gateway.planner_inputs) == 1


def test_refining_preview_keeps_original_export_tail_and_selection(tmp_path):
    gateway = QueueGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        chat, cases = asyncio.run(seed(app.state.store, gateway))
        initial, original, _ = begin(client, app, gateway, chat, cases)
        gateway.plans.append({'title': '调整当前预览', 'steps': [
            {'capability': 'artifact_edit', 'instruction': '将原选择用例的前置条件改为已登录且有权限。'}]})
        gateway.calls.insert(0, ('modify_artifact_tool', {'artifact_id': cases['id'],
            'new_values': {'preconditions': '已登录且有权限'}}))
        refined, _ = post(client, chat, 'refine', '前置条件应该改成已登录且有权限', reply_to=initial['pending'][0]['id'])
        assert refined['status'] == 'needs_confirmation', refined
        state = client.get(f"/api/chats/{chat['id']}/plans/{original['plan_id']}").json()
        assert state['status'] == 'cancelled'
        result, _ = save_workspace(client, refined['pending'][0])
        files = wait_files(client, chat, result)
        sheet = load_workbook(io.BytesIO(client.get(files[0]['url']).content)).active
        rows = list(sheet.values)
        assert len(rows) == 2 and '已登录且有权限' in rows[1]
        assert app.state.store.get('artifact', cases['id'])['items'][1] == cases['items'][1]


def test_restart_after_approval_commit_resumes_tail_using_durable_approval_receipt(tmp_path):
    gateway = QueueGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        chat, cases = asyncio.run(seed(app.state.store, gateway))
        initial, part, _ = begin(client, app, gateway, chat, cases)
        prompt = initial['pending'][0]
        # Deliberately stop between the business approval commit and queue continuation.
        response = client.get('/api/artifacts/' + cases['id'] + '/workspace-grid',
                              params={'proposal_id': prompt['proposal_id']})
        assert response.status_code == 200
        grid = response.json()
        body = TableSave(expected_revision=grid['artifact_revision'], items=grid['proposed_items'],
            layout=grid['layout'], proposal_id=prompt['proposal_id'], prompt_id=prompt['id'],
            profile_id=grid.get('profile_id'), profile_revision=grid.get('profile_revision'),
            client_request_id='approval-window')
        client.portal.call(save_review, app.state.store, app.state.business, app.state.engine, cases['id'], body)
        assert app.state.store.get('artifact', cases['id'])['revision'] == cases['revision'] + 1
        assert not app.state.store.list('frozen_export', chat_id=chat['id'])
    restarted = create_app(tmp_path, gateway)
    with TestClient(restarted) as client:
        files = wait_files(client, chat, {})
        assert len(files) == 1
        assert len(gateway.planner_inputs) == 1
        state = client.get(f"/api/chats/{chat['id']}/plans/{part['plan_id']}").json()
        assert state['status'] == 'completed'
        assert len(restarted.state.store.list('frozen_export', chat_id=chat['id'])) == 1


def test_uncertain_effect_after_crash_is_not_replayed_by_retry(tmp_path):
    gateway = QueueGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        chat, cases = asyncio.run(seed(app.state.store, gateway))
        initial, part, _ = begin(client, app, gateway, chat, cases)
        queued = app.state.store.get('execution_plan', part['plan_id'])
        queued.update(status='running')
        queued['steps'][0].update(status='running', receipts=[])
        app.state.store.put('execution_plan', queued)
    restarted = create_app(tmp_path, gateway)
    with TestClient(restarted) as client:
        state = client.get(f"/api/chats/{chat['id']}/plans/{part['plan_id']}").json()
        assert state['status'] == 'blocked'
        result, _ = post(client, chat, 'retry-uncertain', '重试未完成步骤')
        assert result['status'] == 'failed'
        assert '不能自动重放' in result['message']
        assert len(gateway.planner_inputs) == 1
        assert gateway.executed == ['modify_artifact_tool']
        assert not restarted.state.store.list('frozen_export', chat_id=chat['id'])
