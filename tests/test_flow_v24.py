from fastapi.testclient import TestClient
from tcg.main import create_app
from test_backend_api import Model, setup_chat, start, until

class FlowModel(Model):
    async def generate(self, task, context):
        result = await super().generate(task, context)
        if task == 'analyze_requirement':
            result['report']['summary']='理解登录规则'
            result['report']['strategy']={'depth':'standard','rationale':'覆盖正常和错误凭证'}
            result['report']['diagrams']=[{'title':'业务流程','mermaid':'flowchart TD\n A[登录] --> B[验证凭证]'}]
            if context.get('clarification'):result['report']['questions']=[]
        if task == 'generate_scenarios':
            for item in result['items']:item['requirement_ids']=[i['id'] for i in context['analysis']]
        return result

def test_hitp_gates_and_latest_scenario(tmp_path):
    model=FlowModel(questions=['允许哪些角色登录？'])
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client)
        run=until(client,start(client,chat,experience='reliable',mode='hitp'))
        assert run['status']=='waiting' and run['interrupt']['type']=='clarification',run
        assert not any(t=='generate_cases' for t,_ in model.calls)
        client.post('/api/runs/'+run['id']+'/resume',json={'answer':'注册用户'})
        run=until(client,run)
        assert run['interrupt']['type']=='strategy_review',run
        client.post('/api/runs/'+run['id']+'/resume',json={'approved':True})
        run=until(client,run)
        assert run['interrupt']['type']=='scenario_review',run
        artifact=client.get('/api/artifacts/'+run['interrupt']['artifact_id']).json()
        artifact['items'][0]['title']='人工修改后的场景'
        assert client.put('/api/artifacts/'+artifact['id'],json={'expected_revision':artifact['revision'],'items':artifact['items']}).status_code==200
        client.post('/api/runs/'+run['id']+'/resume',json={'approved':True})
        run=until(client,run)
        assert run['status']=='completed',run
        cases=[c for t,c in model.calls if t=='generate_cases']
        assert cases[0]['scenarios'][0]['title']=='人工修改后的场景'
        assert sum(t=='review_cases' for t,_ in model.calls)==1

def test_auto_keeps_stages_and_no_pauses(tmp_path):
    model=FlowModel(questions=['未说明错误提示文案'])
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client, 'Valid credentials log in.\n\n' * 100)
        run=until(client,start(client,chat,experience='reliable',mode='auto'))
        assert run['status']=='completed',run
        assert [t for t,_ in model.calls]==['analyze_requirement','generate_scenarios','generate_cases','review_cases']
        visible=client.get('/api/chats/'+chat['id']).json()['messages']
        assert len([m for m in visible if m['role']=='assistant'])==1
        assert client.get('/api/artifacts/'+run['artifact_ids'][0]+'/export').status_code==200


def test_profile_override_and_excel_columns(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    model=FlowModel()
    with TestClient(create_app(tmp_path,model)) as client:
        project,chat,_=setup_chat(client)
        profiles=client.get('/api/projects/'+project['id']+'/profiles').json()
        config={'scenario_level':'deep','case_level':'quick','sheet_name':'验收模板','excel_columns':[{'field':'title','header':'用例名称'},{'field':'expected','header':'预期结果'}]}
        run=until(client,start(client,chat,experience='reliable',mode='auto',profile_id=profiles[0]['id'],profile_override=config))
        assert run['status']=='completed',run
        exported=client.get('/api/artifacts/'+run['artifact_ids'][0]+'/export')
        assert exported.status_code==200
        sheet=load_workbook(BytesIO(exported.content)).active
        assert sheet.title=='验收模板'
        assert list(next(sheet.values))==['用例名称','预期结果']
        assert sheet.cell(2,1).value=='Reviewed login case'
        assert client.get('/api/projects/'+project['id']+'/profiles').json()==profiles
