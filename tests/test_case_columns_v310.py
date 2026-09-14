"""Case data changes precede explicit Profile approval and exact frozen exports."""
import base64
import copy
import io
from contextlib import asynccontextmanager
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from tcg.documents import parse_text
from tcg.native_business import NativeBusiness
from tcg.profile_changes import apply_profile_change, preview_profile_change
from tcg.schemas import DomainError
from tcg.storage import Store, dump, now
from tcg.tool_registry import build_tools


@pytest.fixture
def setup(tmp_path):
    store = Store(tmp_path)
    project = store.list('project')[0]
    chat = store.create_chat(project['id'], '用例改列并导出')
    profile = store.list('profile')[0]
    text, chunks = parse_text('用户能登录。')
    source = store.add_source(chat['id'], '登录需求', 'primary', text, chunks)
    case = {'id': 'C1', 'title': '登录', 'scenario_id': '', 'type': 'Business', 'priority': 'P1',
            'preconditions': '', 'steps': [{'action': '登录', 'expected': '登录成功'}],
            'refs': [source['id'] + '#P1'], 'obsolete': '删除前的真实值'}
    artifact = store.put('artifact', {'id': 'cases', 'chat_id': chat['id'], 'project_id': project['id'],
        'type': 'cases', 'title': '测试用例', 'revision': 1, '_visible': True, '_profile': profile['config'],
        '_source_ids': [source['id']], '_source_roles': {source['id']: 'primary'}, 'items': [case], 'report': {}})
    store.db.execute('INSERT INTO revisions VALUES(?,?,?,?,?,?)',
        ('cases', 1, dump(artifact), now(), 'test_fixture', '{}'))
    class Gateway:
        async def generate_native(self, task, context, schema, instruction):
            rows = copy.deepcopy(context['items'])
            for row in rows:
                row['test_data'] = '当前需求中的登录用户'
                row['title'] = '模型不得顺便修改标题'
            return {'items': rows, 'report': {'summary': '已补充测试数据'}}
    class Pipeline:
        @asynccontextmanager
        async def edit_session(self, chat_id):
            yield
        async def on_artifact_changed(self, artifact):
            pass
    business = NativeBusiness(store, Gateway())
    def tools(prompt=None):
        return {t.name: t for t in build_tools(store, business, Pipeline(), chat,
            {'content': '给用例新增测试数据和执行人，删除废弃列，然后导出', 'reply_to': (prompt or {}).get('id')}, prompt)}
    yield SimpleNamespace(store=store, chat=chat, profile=profile, artifact=artifact, tools=tools, business=business)
    store.close()


async def propose_and_apply_cases(c, **kwargs):
    result = await c.tools()['modify_case_columns_tool'].ainvoke({'artifact_id': 'cases',
        'upsert_columns': [{'field': 'tester', 'header': '执行人'}], 'export_after_approval': True, **kwargs})
    assert result['status'] == 'needs_confirmation', result
    assert c.store.get('artifact', 'cases')['revision'] == 1
    assert c.store.get('profile', c.profile['id'])['version'] == 1
    applied = await c.tools(result['pending'][0])['apply_artifact_preview_tool'].ainvoke({})
    assert applied['status'] == 'needs_confirmation', applied
    assert c.store.get('artifact', 'cases')['revision'] == 2
    assert c.store.get('profile', c.profile['id'])['version'] == 1
    return applied['pending'][0]


@pytest.mark.asyncio
async def test_case_columns_then_profile_then_export_contains_exact_new_revision(setup):
    c = setup
    p = await propose_and_apply_cases(c, upsert_columns=[{'field': 'test_data', 'header': '测试数据'},
        {'field': 'tester', 'header': '执行人'}], remove_columns=['obsolete'])
    current = c.store.get('artifact', 'cases')['items'][0]
    assert current['title'] == '登录'
    assert current['test_data'] == '当前需求中的登录用户'
    assert current['tester'] == '' and 'obsolete' not in current
    assert c.store.revision('cases', 1)['items'][0]['obsolete'] == '删除前的真实值'
    assert not c.store.list('frozen_export')
    preview = preview_profile_change(c.store, c.chat['id'], p['id'])
    applied = apply_profile_change(c.store, c.chat['id'], p['id'], 1, ['excel_columns'], write_messages=True)
    url = applied['parts'][0]['files'][0]['url']
    exported = c.store.get('frozen_export', url.split('/')[-1])
    assert exported['revision'] == 2 and exported['profile_version'] == 2
    sheet = load_workbook(io.BytesIO(base64.b64decode(exported['_bytes']))).active
    rows = list(sheet.values)
    assert rows[1][rows[0].index('测试数据')] == '当前需求中的登录用户'
    assert rows[1][rows[0].index('执行人')] is None
    assert len(c.store.list('frozen_export')) == 1
    messages = c.store.list('message', chat_id=c.chat['id'])
    assistant = [m for m in messages if m['role'] == 'assistant'][-1]
    assert assistant['metadata']['turn_response']['parts'] == applied['parts']
    with pytest.raises(DomainError):
        apply_profile_change(c.store, c.chat['id'], p['id'], 1)
    assert len(c.store.list('frozen_export')) == 1


@pytest.mark.asyncio
async def test_stale_case_revision_blocks_profile_and_deferred_export(setup):
    c = setup
    p = await propose_and_apply_cases(c)
    current = c.store.get('artifact', 'cases')
    await c.business.revise(current, new_values={'title': '新的手动修改'}, preview=False)
    with pytest.raises(DomainError, match='用例已改变'):
        apply_profile_change(c.store, c.chat['id'], p['id'], 1)
    assert c.store.get('profile', c.profile['id'])['version'] == 1
    assert not c.store.list('frozen_export')


@pytest.mark.asyncio
async def test_discard_and_protected_fields_never_export(setup):
    c = setup
    protected = await c.tools()['modify_case_columns_tool'].ainvoke({'artifact_id': 'cases', 'remove_columns': ['steps']})
    assert protected['status'] == 'needs_input'
    assert '核心' in protected['message']
    p = await propose_and_apply_cases(c)
    result = await c.tools(p)['discard_template_tool'].ainvoke({})
    assert result['status'] == 'succeeded'
    assert not c.store.list('frozen_export')
    assert c.store.get('profile', c.profile['id'])['version'] == 1


@pytest.mark.asyncio
async def test_export_called_after_preview_waits_and_remembers_request(setup):
    c = setup
    first = await c.tools()['modify_case_columns_tool'].ainvoke({'artifact_id': 'cases',
        'upsert_columns': [{'field': 'tester', 'header': '执行人'}]})
    queued = await c.tools(first['pending'][0])['export_artifact_tool'].ainvoke({'artifact_ids': ['cases']})
    assert queued['status'] == 'needs_confirmation'
    assert not c.store.list('frozen_export')
    adopted = await c.tools(first['pending'][0])['apply_artifact_preview_tool'].ainvoke({})
    p = adopted['pending'][0]
    result = await c.tools(p)['apply_profile_tool'].ainvoke({})
    assert result['parts'][0]['type'] == 'files'
    assert len(c.store.list('frozen_export')) == 1


@pytest.mark.asyncio
async def test_manual_column_change_prepares_profile_without_export(setup):
    from tcg.case_columns import attach_manual_column_sync
    c = setup
    before = c.store.get('artifact', 'cases')
    saved = await c.business.revise(before, new_values={'test_data': '人工输入'}, preview=False)
    result = attach_manual_column_sync(c.store, saved, before,
        {'added': [{'field': 'test_data', 'header': '测试数据'}], 'removed': []})
    assert result['status'] == 'needs_confirmation'
    p = result['pending'][0]
    applied = apply_profile_change(c.store, c.chat['id'], p['id'], 1)
    assert applied['parts'] == []
    assert not c.store.list('frozen_export')
    assert c.store.get('artifact', 'cases')['items'][0]['test_data'] == '人工输入'


@pytest.mark.asyncio
async def test_core_export_column_can_be_hidden_without_deleting_steps(setup):
    c = setup
    result = await c.tools()['modify_case_columns_tool'].ainvoke({'artifact_id': 'cases',
        'hide_columns': ['steps'], 'export_after_approval': True})
    assert result['status'] == 'needs_confirmation', result
    applied = await c.tools(result['pending'][0])['apply_artifact_preview_tool'].ainvoke({})
    current = c.store.get('artifact', 'cases')
    assert current['items'][0]['steps'] == c.artifact['items'][0]['steps']
    p = applied['pending'][0]
    applied = apply_profile_change(c.store, c.chat['id'], p['id'], 1)
    assert not any(col['field'] == 'steps' for col in applied['profile']['config']['excel_columns'])
    assert applied['parts'][0]['type'] == 'files'


@pytest.mark.parametrize('change', [{'added': ['bad']}, {'added': [{}]}, {'removed': 'tester'},
    {'added': [{'header': '缺字段名'}]}, {'extra': True}])
def test_malformed_manual_column_changes_are_domain_errors(setup, change):
    from tcg.case_columns import manual_column_plan
    with pytest.raises(DomainError):
        manual_column_plan(setup.store, setup.artifact, change)


@pytest.mark.asyncio
async def test_amending_profile_columns_rebinds_deferred_export_and_case_guard(setup):
    c = setup
    old = await propose_and_apply_cases(c)
    request_id = old['deferred_export_id']
    changed = await c.tools(old)['modify_profile_tool'].ainvoke({
        'upsert_columns': [{'field': 'tester', 'header': '测试执行人'}]})
    assert changed['status'] == 'needs_confirmation', changed
    pending = changed['pending'][0]
    assert pending['case_binding']['artifact_id'] == old['case_binding']['artifact_id']
    assert pending['case_binding']['revision'] == 2
    assert pending['deferred_export_id'] == request_id
    request = c.store.get('deferred_export', request_id)
    assert request['prompt_id'] == pending['id'] and request['template_ids'] == pending['template_ids']
    assert next(col for col in request['column_config'] if col['field'] == 'tester')['header'] == '测试执行人'
    stale = await c.tools(old)['apply_profile_tool'].ainvoke({})
    assert stale['status'] == 'needs_input'
    assert not c.store.list('frozen_export')
    applied = await c.tools(pending)['apply_profile_tool'].ainvoke({})
    assert applied['status'] == 'succeeded'
    exported = c.store.get('frozen_export', applied['parts'][0]['files'][0]['url'].split('/')[-1])
    assert exported['revision'] == 2
    sheet = load_workbook(io.BytesIO(base64.b64decode(exported['_bytes']))).active
    assert '测试执行人' in list(sheet.values)[0]
    assert len(c.store.list('frozen_export')) == 1


@pytest.mark.asyncio
async def test_amended_profile_still_refuses_newer_case_revision(setup):
    c = setup
    old = await propose_and_apply_cases(c)
    changed = await c.tools(old)['modify_profile_tool'].ainvoke({
        'upsert_columns': [{'field': 'tester', 'header': '测试执行人'}]})
    pending = changed['pending'][0]
    await c.business.revise(c.store.get('artifact', 'cases'), new_values={'title': '另一次修改'}, preview=False)
    with pytest.raises(DomainError, match='用例已改变'):
        apply_profile_change(c.store, c.chat['id'], pending['id'], 1)
    assert not c.store.list('frozen_export')
