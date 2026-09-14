"""Human table resolutions stay version-bound and publish through native gates."""
import copy
import io

import pytest
from openpyxl import load_workbook

from tcg.native_views import current_prompt
from tcg.pipeline import PipelineRuntime
from tcg.schemas import DomainError
from tcg.storage import uid
from tcg.table_review import TableDraft, TableSave, project_draft, review_view, save_review
from test_review_proposals_v310 import setup, review_gate
from test_native_pipeline_v300 import settled


async def opened(store, business, runtime, chat):
    run = await review_gate(runtime, chat)
    view = await review_view(store, runtime, run['current_artifact_id'])
    body = {'items': copy.deepcopy(view['proposed_items']), 'expected_revision': view['artifact_revision'],
        'profile_id': view['profile_id'], 'profile_revision': view['profile_revision'],
        'run_id': view['run_id'], 'proposal_id': view['proposal_id'], 'prompt_id': view['prompt_id']}
    return run, view, body


@pytest.mark.asyncio
async def test_mixed_rejection_manual_steps_commits_once_through_gate(setup):
    store, chat, model, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run, view, body = await opened(store, business, runtime, chat)
        aid = view['artifact_id']
        body['items'][0]['title'] = view['original_items'][0]['title']
        body['items'][0]['steps'][0] = {'action': '用户登录后点击入口', 'expected': '显示登录后的主界面'}
        request = TableSave(**body, client_request_id='resolution-one')
        before = store.get('artifact', aid)
        result = await save_review(store, business, runtime, aid, request)
        assert result['run']['status'] == 'queued'
        assert store.get('artifact', aid) == before
        # A transport retry is a receipt read even while the pipeline is queued.
        await save_review(store, business, runtime, aid, request)
        done = await settled(runtime, run['id'])
        assert done['status'] == 'completed', done.get('error')
        after = store.get('artifact', aid)
        assert after['revision'] == 2
        assert after['items'][0]['title'] == view['original_items'][0]['title']
        assert after['items'][0]['steps'][0] == body['items'][0]['steps'][0]
        assert after['items'][0]['refs']
        assert len(store.list('table_review_resolution')) == 1
        receipt = store.get('native_approval_receipt', 'approval:' + body['prompt_id'])
        assert receipt['parts'][0]['revision'] == 2
        await save_review(store, business, runtime, aid, request)
        assert store.get('artifact', aid)['revision'] == 2
        assert len([t for t, _ in model.calls if t == 'review_cases']) == 1
    finally:
        await runtime.stop()


@pytest.mark.asyncio
async def test_export_preview_matches_projection_and_does_not_accept_gate(setup):
    store, chat, model, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run, view, body = await opened(store, business, runtime, chat)
        aid = view['artifact_id']
        body['items'][0]['preconditions'] = '人工草稿，已登录'
        body['items'][0]['steps'].append({'action': '刷新', 'expected': '继续显示页面'})
        for layout in ('case', 'step'):
            request = TableDraft(**body, layout=layout)
            grid = project_draft(store, business, aid, request)
            workbook = load_workbook(io.BytesIO(project_draft(store, business, aid, request, export=True)))
            rows = list(workbook.active.values)
            assert list(rows[0]) == [c['header'] for c in grid['columns']]
            assert [[str(c) if c is not None else '' for c in row] for row in rows[1:]] == [r['cells'] for r in grid['rows']]
        assert store.get('artifact', aid)['revision'] == 1
        assert store.run(run['id'])['status'] == 'waiting'
        assert store.list('table_review_resolution') == []
    finally:
        await runtime.stop()


@pytest.mark.asyncio
@pytest.mark.parametrize('break_binding', ['profile', 'artifact', 'prompt', 'references', 'steps'])
async def test_stale_or_invalid_drafts_leave_artifact_sources_and_gate_untouched(setup, break_binding):
    store, chat, model, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run, view, body = await opened(store, business, runtime, chat)
        aid = view['artifact_id']
        if break_binding == 'profile':
            profile = store.get('profile', body['profile_id'])
            store.update_profile(profile['id'], profile['name'], profile['config'], profile['version'])
        elif break_binding == 'artifact':
            old = store.get('artifact', aid)
            store.revise_artifact(aid, 1, old['items'])
        elif break_binding == 'prompt':
            body['prompt_id'] = 'expired'
        elif break_binding == 'references':
            body['items'][0]['refs'] = ['forged-ref']
        else:
            body['items'][0]['steps'] = [{'action': 'missing expected'}]
        artifact = store.get('artifact', aid)
        sources = store.list('source', chat_id=chat['id'])
        with pytest.raises(DomainError):
            await save_review(store, business, runtime, aid, TableSave(**body, client_request_id='invalid'))
        assert store.get('artifact', aid) == artifact
        assert store.list('source', chat_id=chat['id']) == sources
        assert store.list('table_review_resolution') == []
        assert store.run(run['id'])['status'] == 'waiting'
    finally:
        await runtime.stop()


@pytest.mark.asyncio
async def test_artifact_preview_saves_human_resolution_and_approval_receipt(setup):
    store, chat, model, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run = await runtime.start_run(chat['id'], {'mode': 'auto'})
        run = await settled(runtime, run['id'])
        artifact = store.get('artifact', run['current_artifact_id'])
        proposed = await business.revise(artifact, new_values={'preconditions': 'AI 建议：已登录'}, preview=True)
        proposal = {**proposed, 'id': uid('revprop_'), 'chat_id': chat['id'], 'project_id': chat['project_id']}
        store.put('native_revision_proposal', proposal)
        prompt = {'id': 'revision:' + proposal['id'], 'kind': 'artifact_proposal', 'artifact_id': artifact['id'],
                  'artifact_revision': artifact['revision'], 'proposal_id': proposal['id']}
        store.put('chat', {**store.get('chat', chat['id']), '_native_artifact_prompt': prompt})
        view = await review_view(store, runtime, artifact['id'])
        assert view['proposal_kind'] == 'native_revision_proposal'
        rows = copy.deepcopy(view['proposed_items'])
        rows[0]['preconditions'] = '人工确认：使用已认证的有效账号'
        body = TableSave(items=rows, expected_revision=view['artifact_revision'], profile_id=view['profile_id'],
            profile_revision=view['profile_revision'], proposal_id=view['proposal_id'], prompt_id=view['prompt_id'],
            client_request_id='artifact-resolution')
        result = await save_review(store, business, runtime, artifact['id'], body)
        assert result['artifact']['revision'] == artifact['revision'] + 1
        assert result['artifact']['items'][0]['preconditions'] == rows[0]['preconditions']
        assert store.get('native_revision_proposal', proposal['id'])['_applied']
        assert store.get('native_approval_receipt', 'approval:' + prompt['id'])['parts'][0]['revision'] == artifact['revision'] + 1
        await settled(runtime, run['id'])
        assert (await current_prompt(store, runtime, store.get('chat', chat['id'])))['kind'] == 'case_result_review'
    finally:
        await runtime.stop()


@pytest.mark.asyncio
async def test_history_is_read_only_and_foreign_proposal_is_denied(setup):
    store, chat, model, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run, view, body = await opened(store, business, runtime, chat)
        history = await review_view(store, runtime, view['artifact_id'], revision=1)
        assert history['read_only'] and history['proposal_id'] is None
        with pytest.raises(DomainError, match='历史'):
            await save_review(store, business, runtime, view['artifact_id'],
                TableSave(**body, revision=1, client_request_id='history'))
        proposal = store.get('review_proposal', body['proposal_id'])
        store.put('review_proposal', {**proposal, 'id': 'foreign', 'chat_id': 'another-chat'})
        with pytest.raises(DomainError, match='不属于'):
            await review_view(store, runtime, view['artifact_id'], proposal_id='foreign')
    finally:
        await runtime.stop()


@pytest.mark.asyncio
async def test_historical_projection_and_excel_remain_available_after_source_deactivation(setup):
    store, chat, model, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run = await runtime.start_run(chat['id'], {'mode': 'auto'})
        run = await settled(runtime, run['id'])
        artifact_id = run['current_artifact_id']
        original = store.revision(artifact_id, 1)
        for source_id in original['_source_ids']:
            store.deactivate_source(source_id)
        view = await review_view(store, runtime, artifact_id, revision=1)
        assert view['read_only']
        request = TableDraft(items=view['original_items'], expected_revision=1, revision=1,
            profile_id=view['profile_id'], profile_revision=view['profile_revision'])
        grid = project_draft(store, business, artifact_id, request)
        workbook = load_workbook(io.BytesIO(project_draft(store, business, artifact_id, request, export=True)))
        cells = [[str(value) if value is not None else '' for value in row]
                 for row in list(workbook.active.values)[1:]]
        assert cells == [row['cells'] for row in grid['rows']]
        altered = request.model_copy(deep=True)
        altered.items[0]['title'] = '篡改历史版本'
        for export in (False, True):
            with pytest.raises(DomainError, match='历史版本'):
                project_draft(store, business, artifact_id, altered, export=export)
        live = request.model_copy(update={'revision': None,
            'expected_revision': store.get('artifact', artifact_id)['revision']})
        with pytest.raises(DomainError):
            project_draft(store, business, artifact_id, live)
        assert store.revision(artifact_id, 1) == original
    finally:
        await runtime.stop()
