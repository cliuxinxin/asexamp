"""Focused scenario-template isolation, learning and deterministic export coverage."""
import copy
import io
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tcg.documents import export_artifact
from tcg.main import create_app
from tcg.schemas import DEFAULT_PROFILE, DomainError, profile_config
from tcg.workflow import WorkflowEngine
from test_backend_api import setup_chat, start, until
from test_workflow_v25 import FlowModel


SCENARIO_COLUMNS = [
    {'field': 'description', 'header': '场景说明', 'definition': '说明验证的业务路径'},
    {'field': 'requirement_ids', 'header': '需求编号'},
    {'field': 'refs', 'header': '证据编号'},
]


def current_config():
    return profile_config({
        'language': 'English', 'scope': '当前业务范围', 'additional_rules': '保留写作规则',
        'excel_columns': [{'field': 'title', 'header': '既有用例标题'}],
        'excel_layout': 'step', 'sheet_name': '既有用例', 'filename_pattern': 'cases_{project}.xlsx',
        'scenario_excel_columns': [{'field': 'title', 'header': '既有场景标题'}],
        'scenario_sheet_name': '既有场景', 'scenario_filename_pattern': 'scenarios_{project}.xlsx',
    })


@pytest.mark.parametrize('kinds', [['scenarios'], ['cases'], ['scenarios', 'cases']])
def test_template_learning_updates_only_detected_kinds(kinds):
    current = current_config()
    proposal = profile_config({
        'scenario_excel_columns': SCENARIO_COLUMNS, 'scenario_sheet_name': '学习场景',
        'excel_columns': [{'field': 'description', 'header': '学习用例描述'}], 'sheet_name': '学习用例',
    })
    proposal['template_kinds'] = kinds
    saved = copy.deepcopy(proposal)
    merged, _ = WorkflowEngine.__new__(WorkflowEngine).template_config(current, proposal, kinds)
    assert merged['scenario_excel_columns'] == (SCENARIO_COLUMNS if 'scenarios' in kinds else current['scenario_excel_columns'])
    assert merged['sheet_name'] == ('学习用例' if 'cases' in kinds else current['sheet_name'])
    if kinds == ['scenarios']:
        assert {k: v for k, v in merged.items() if not k.startswith('scenario_')} == {k: v for k, v in current.items() if not k.startswith('scenario_')}
    assert 'template_kinds' not in merged
    assert proposal == saved


def test_empty_unrecognized_template_preserves_custom_layout_and_blanks():
    current = current_config()
    engine = WorkflowEngine.__new__(WorkflowEngine)
    merged, notes = engine.template_config(current, {
        'excel_columns': [], 'excel_layout': 'case', 'sheet_name': 'Test Cases', 'filename_pattern': None,
        'scenario_excel_columns': [], 'scenario_sheet_name': '', 'scenario_filename_pattern': ' ',
    }, ['scenarios', 'cases'])
    for field in ('excel_columns', 'excel_layout', 'sheet_name', 'filename_pattern',
                  'scenario_excel_columns', 'scenario_sheet_name', 'scenario_filename_pattern'):
        assert merged[field] == current[field]
    assert notes


def test_scenario_profile_validation_names_column_and_keeps_case_restrictions():
    assert 'template_kinds' not in profile_config({'template_kinds': ['scenarios']})
    for field in ('refs', 'requirement_ids'):
        profile_config({'scenario_excel_columns': [{'field': field, 'header': '追溯'}]})
        with pytest.raises(DomainError, match='来源或内部记录'):
            profile_config({'excel_columns': [{'field': field, 'header': '追溯'}]})
    for column in ({'field': '_profile', 'header': '禁止'}, {'field': 'title', 'header': '标题', 'definition': 42}):
        with pytest.raises(DomainError, match='scenario_excel_columns'):
            profile_config({'scenario_excel_columns': [column]})


class ScenarioTemplateModel(FlowModel):
    def __init__(self, kinds=None):
        super().__init__()
        self.kinds = kinds or ['scenarios']

    async def generate(self, task, context):
        if task == 'learn_template':
            self.calls.append((task, copy.deepcopy(context)))
            return {'template_kinds': self.kinds, 'config': {
                'scenario_excel_columns': SCENARIO_COLUMNS,
                'scenario_sheet_name': '学习场景', 'scenario_filename_pattern': '场景_{project}_{date}.xlsx',
                'excel_columns': [], 'sheet_name': 'Test Cases',
            }, 'summary': '提取场景列定义并保留需求与证据追溯。'}
        return await super().generate(task, context)


def test_scenario_learning_is_proposal_and_generation_receives_format_references(tmp_path):
    model = ScenarioTemplateModel()
    with TestClient(create_app(tmp_path, model)) as client:
        project, chat, _ = setup_chat(client)
        example = client.post('/api/chats/' + chat['id'] + '/sources/text', json={
            'name': '场景模板定义', 'text': '场景说明：说明验证的业务路径；需求编号；证据编号', 'role': 'example',
        }).json()
        profiles_before = client.get('/api/projects/' + project['id'] + '/profiles').json()
        run = until(client, start(client, chat, experience='reliable', intent='learn_template', profile_override=current_config()))
        assert run['status'] == 'completed', run.get('error')
        proposal = client.get('/api/artifacts/' + run['artifact_ids'][-1]).json()
        config = proposal['report']['config']
        assert proposal['report']['template_kinds'] == ['scenarios']
        assert 'template_kinds' not in config
        assert config['excel_columns'] == current_config()['excel_columns']
        assert config['sheet_name'] == current_config()['sheet_name']
        assert config['scenario_excel_columns'] == SCENARIO_COLUMNS
        assert client.get('/api/projects/' + project['id'] + '/profiles').json() == profiles_before
        generated = until(client, start(client, chat, experience='reliable', intent='generate_scenario', profile_override=config))
        assert generated['status'] == 'completed', generated.get('error')
        context = next(context for task, context in model.calls if task == 'generate_scenarios')
        assert context['profile']['scenario_excel_columns'] == SCENARIO_COLUMNS
        assert {row['source_id'] for row in context['format_references']} == {example['id']}
        assert all(row['role'] != 'example' for row in context['evidence'])
        assert 'scenario_excel_columns' in context['format_instruction']


def test_scenario_export_api_uses_legacy_defaults_or_selected_profile_without_case_checks(tmp_path, monkeypatch):
    model = FlowModel()
    with TestClient(create_app(tmp_path, model)) as client:
        project, chat, _ = setup_chat(client)
        run = until(client, start(client, chat, experience='reliable', intent='generate_scenario'))
        assert run['status'] == 'completed', run.get('error')
        aid = run['artifact_ids'][-1]
        stored = client.app.state.store.get('artifact', aid)
        stored['_profile'] = {'excel_columns': [{'field': 'missing_case_field', 'header': '必填用例字段'}]}
        client.app.state.store.put('artifact', stored)
        columns = [{'field': 'refs', 'header': '=证据'}, {'field': 'title', 'header': '场景标题'}]
        profile = client.post('/api/projects/' + project['id'] + '/profiles', json={'name': '场景导出', 'config': {
            'scenario_excel_columns': columns, 'scenario_sheet_name': '场景/[清单]',
            'scenario_filename_pattern': '已选场景_{project}.xlsx',
            'excel_columns': [{'field': 'missing_case_field', 'header': '必填用例字段'}],
        }}).json()
        marker = len(model.calls)
        def forbidden_case_check(*args, **kwargs):
            raise AssertionError('Scenario export must not invoke case completion checks')
        monkeypatch.setattr('tcg.case_fields.template_check', forbidden_case_check)
        options = client.get('/api/artifacts/' + aid + '/export-options')
        assert options.status_code == 200, options.text
        assert options.json()['kind'] == 'scenarios'
        assert options.json()['snapshot']['scenario_sheet_name'] == 'Test Scenarios'
        default = client.get('/api/artifacts/' + aid + '/export')
        assert default.status_code == 200, default.text
        assert '_scenarios_' in unquote(default.headers['content-disposition'])
        assert load_workbook(io.BytesIO(default.content)).active.title == 'Test Scenarios'
        custom = client.get('/api/artifacts/' + aid + '/export', params={'profile_id': profile['id'], 'ids': stored['items'][0]['id']})
        assert custom.status_code == 200, custom.text
        assert '已选场景_' in unquote(custom.headers['content-disposition'])
        sheet = load_workbook(io.BytesIO(custom.content)).active
        assert sheet.title == '场景__清单_'
        assert list(next(sheet.values)) == ["'=证据", '场景标题']
        assert sheet.max_row == 2 and sheet.cell(2, 1).value == '\n'.join(stored['items'][0]['refs'])
        for selected in ('', 'missing'):
            assert client.get('/api/artifacts/' + aid + '/export', params={'ids': selected}).status_code == 400
        assert client.post('/api/artifacts/' + aid + '/complete-fields', json={}).status_code == 400
        assert len(model.calls) == marker
        assert client.app.state.store.get('artifact', aid) == stored


def test_scenario_export_serializes_custom_nested_values_and_validates_selection():
    artifact = {'type': 'scenarios', '_profile': {'scenario_excel_columns': [
        {'field': 'payload', 'header': '补充信息'}, {'field': 'refs', 'header': '证据'},
    ]}, 'items': [{'id': 'S1', 'payload': [{'path': '登录'}], 'refs': ['=danger', 'src#P1']}]}
    sheet = load_workbook(io.BytesIO(export_artifact(artifact))).active
    assert sheet.cell(2, 1).value == '{"path": "登录"}'
    assert sheet.cell(2, 2).value == "'=danger\nsrc#P1"
    assert sheet.cell(2, 2).alignment.wrap_text is True
    assert sheet.freeze_panes == 'A2'
    for selected in ([], ['missing']):
        with pytest.raises(DomainError, match='ID 无效'):
            export_artifact(artifact, selected=selected)
