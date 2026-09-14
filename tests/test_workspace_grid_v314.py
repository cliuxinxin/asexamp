"""One workspace projects and saves every native artifact type."""
import copy
import io

import pytest
from openpyxl import load_workbook

from tcg.pipeline import PipelineRuntime
from tcg.table_projection import artifact_table_projection
from tcg.table_review import TableDraft, TableSave, project_draft, review_view, save_review
from test_manual_tables_v310 import scenarios_ready
from test_native_journey_v300 import native_journey
from test_native_pipeline_v300 import settled
from test_review_proposals_v310 import review_gate, setup


def test_analysis_and_scenario_projection_keep_stable_rows_and_profile_columns():
    analysis = {'type': 'analysis', 'items': [
        {'id': 'REQ-1', 'title': '登录', 'description': '用户可以登录。', 'refs': ['ev-1']},
    ]}
    projected = artifact_table_projection(analysis)
    assert projected == {
        'columns': [
            {'field': 'id', 'header': 'Requirement ID'},
            {'field': 'title', 'header': 'Title'},
            {'field': 'description', 'header': 'Description'},
        ],
        'rows': [{'item_id': 'REQ-1', 'step_index': None,
                  'cells': ['REQ-1', '登录', '用户可以登录。']}],
        'layout': 'case',
    }

    scenarios = {'type': 'scenarios', '_profile': {'scenario_excel_columns': [
        {'field': 'title', 'header': '场景'},
        {'field': 'requirement_ids', 'header': '需求编号'},
    ]}, 'items': [
        {'id': 'SC-1', 'title': '密码登录', 'requirement_ids': ['REQ-1'], 'refs': ['ev-1']},
        {'id': 'SC-2', 'title': '=危险标题', 'requirement_ids': ['REQ-2'], 'refs': ['ev-2']},
    ]}
    before = copy.deepcopy(scenarios)
    projected = artifact_table_projection(scenarios, selected=['SC-2'])
    assert projected['columns'] == scenarios['_profile']['scenario_excel_columns']
    assert projected['rows'] == [
        {'item_id': 'SC-2', 'step_index': None, 'cells': ["'=危险标题", 'REQ-2']},
    ]
    assert scenarios == before


def test_scenario_manual_save_and_history_round_trip_through_workspace(native_journey):
    j = native_journey
    analysis, scenarios = scenarios_ready(j)
    base = '/api/artifacts/' + scenarios['id'] + '/workspace-grid'
    opened = j.client.get(base)
    assert opened.status_code == 200, opened.text
    view = opened.json()
    assert view['artifact_type'] == 'scenarios'
    assert view['mode'] == 'manual'
    assert not view['read_only'] and view['proposal_id'] is None

    rows = copy.deepcopy(view['proposed_items'])
    rows.append({'id': 'SC-MANUAL', 'title': '人工补充并发登录',
        'description': '两个设备同时登录，核对访问结果。', 'priority': 'P2',
        'requirement_ids': [], 'refs': []})
    report = copy.deepcopy(scenarios['report'])
    report['summary'] = '人工核对后的场景范围'
    body = {'items': rows, 'report': report, 'expected_revision': scenarios['revision'],
            'client_request_id': 'workspace-scenario-manual'}
    saved = j.client.post(base + '/save', json=body)
    assert saved.status_code == 200, saved.text
    result = saved.json()['artifact']
    assert result['revision'] == scenarios['revision'] + 1
    assert result['report']['summary'] == report['summary']
    manual = next(row for row in result['items'] if row['id'] == 'SC-MANUAL')
    assert manual['requirement_ids'] == [] and manual['refs']
    assert manual['_independent_origin']['source_id']
    assert j.app.state.store.get('artifact', analysis['id']) == analysis

    repeated = j.client.post(base + '/save', json=body)
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()['artifact']['revision'] == result['revision']

    history = j.client.get(base, params={'revision': scenarios['revision']})
    assert history.status_code == 200, history.text
    historical = history.json()
    assert historical['mode'] == 'read_only' and historical['read_only']
    assert historical['original_items'] == scenarios['items']
    projected = j.client.post(base + '/project', json={
        'items': historical['original_items'], 'expected_revision': scenarios['revision'],
        'revision': scenarios['revision']})
    assert projected.status_code == 200, projected.text
    assert projected.json()['rows'] == historical['original_rows']


def test_analysis_history_exports_exact_projection_and_rejects_column_edits(native_journey):
    j = native_journey
    analysis, _ = scenarios_ready(j)
    body = TableDraft(items=copy.deepcopy(analysis['items']),
                      expected_revision=analysis['revision'], revision=analysis['revision'])
    grid = project_draft(j.app.state.store, j.app.state.business, analysis['id'], body)
    workbook = load_workbook(io.BytesIO(project_draft(
        j.app.state.store, j.app.state.business, analysis['id'], body, export=True)))
    values = [[value if value is not None else '' for value in row]
              for row in workbook.active.values]
    assert values == [[column['header'] for column in grid['columns']]] + [
        row['cells'] for row in grid['rows']]
    changed = body.model_copy(update={'column_changes': {
        'added': [{'field': 'owner', 'header': '负责人'}], 'removed': []}})
    response = j.client.post('/api/artifacts/' + analysis['id'] + '/workspace-grid/project',
                             json=changed.model_dump())
    assert response.status_code == 409, response.text


def test_historical_scenario_defaults_to_frozen_profile_after_live_profile_changes(native_journey):
    j = native_journey
    _, scenarios = scenarios_ready(j)
    base = '/api/artifacts/' + scenarios['id'] + '/workspace-grid'
    before = j.client.get(base, params={'revision': scenarios['revision']})
    assert before.status_code == 200, before.text
    frozen = before.json()
    profile = j.app.state.store.list('profile', project_id=j.project['id'])[0]
    config = copy.deepcopy(profile['config'])
    config['scenario_excel_columns'] = [{'field': 'title', 'header': 'NEW LIVE HEADER'}]
    j.app.state.store.update_profile(profile['id'], profile['name'], config, profile['version'])

    after = j.client.get(base, params={'revision': scenarios['revision']})
    assert after.status_code == 200, after.text
    historical = after.json()
    assert historical['columns'] == frozen['columns']
    assert historical['original_rows'] == frozen['original_rows']
    assert historical['profile_id'] is None and historical['profile_revision'] is None
    export = j.client.post(base + '/export', json={
        'items': historical['original_items'], 'expected_revision': scenarios['revision'],
        'revision': scenarios['revision']})
    assert export.status_code == 200, export.text
    headers = list(load_workbook(io.BytesIO(export.content)).active.values)[0]
    assert list(headers) == [column['header'] for column in frozen['columns']]


def test_explicit_reject_discards_metadata_only_column_proposal_idempotently(native_journey):
    j = native_journey
    j.turn('生成测试用例并评审', 'start_pipeline_tool', {'mode': 'auto'}, mode='auto')
    j.completed()
    store = j.app.state.store
    cases = next(value for value in store.list('artifact', chat_id=j.chat['id'])
                 if value['type'] == 'cases')
    profile = store.list('profile', project_id=j.project['id'])[0]
    sources = store.list('source', chat_id=j.chat['id'])
    prepared = j.turn('隐藏步骤导出列', 'modify_case_columns_tool', {
        'artifact_id': cases['id'], 'hide_columns': ['steps']}, status='needs_confirmation')
    prompt = prepared['pending'][0]
    proposal = store.get('artifact_proposal', prompt['proposal_id'])
    assert proposal['changes'] == []

    base = '/api/artifacts/' + cases['id'] + '/workspace-grid'
    opened = j.client.get(base, params={'proposal_id': proposal['id']})
    assert opened.status_code == 200, opened.text
    view = opened.json()
    body = {'items': view['original_items'],
            'expected_revision': view['artifact_revision'], 'profile_id': view['profile_id'],
            'profile_revision': view['profile_revision'], 'proposal_id': view['proposal_id'],
            'prompt_id': view['prompt_id'], 'proposal_decision': 'reject',
            'client_request_id': 'reject-metadata-only-column-proposal'}
    response = j.client.post(base + '/save', json=body)
    assert response.status_code == 200, response.text
    assert response.json()['rejected'] is True
    assert store.get('artifact', cases['id']) == cases
    assert store.get('profile', profile['id']) == profile
    assert store.list('source', chat_id=j.chat['id']) == sources
    assert store.get('artifact_proposal', proposal['id'])['status'] == 'rejected'
    assert store.get('chat', j.chat['id']).get('_native_artifact_prompt') is None

    repeated = j.client.post(base + '/save', json=body)
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()['rejected'] is True
    assert store.get('artifact', cases['id']) == cases
    assert store.get('profile', profile['id']) == profile


@pytest.mark.asyncio
async def test_rejected_review_receipt_replays_rejected_action_after_interrupted_schedule(setup):
    store, chat, _, business = setup
    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        run = await review_gate(runtime, chat)
        view = await review_view(store, runtime, run['current_artifact_id'])
        body = TableSave(items=copy.deepcopy(view['original_items']),
            expected_revision=view['artifact_revision'], profile_id=view['profile_id'],
            profile_revision=view['profile_revision'], run_id=view['run_id'],
            proposal_id=view['proposal_id'], prompt_id=view['prompt_id'],
            client_request_id='rejected-restart-retry')
        runtime._schedule = lambda *args, **kwargs: None
        first = await save_review(store, business, runtime, view['artifact_id'], body)
        assert first['rejected'] and store.get('artifact', view['artifact_id'])['revision'] == 1
    finally:
        await runtime.stop()

    runtime = PipelineRuntime(store, business)
    await runtime.start()
    try:
        repeated = await save_review(store, business, runtime, view['artifact_id'], body)
        assert repeated['rejected']
        done = await settled(runtime, run['id'])
        assert done['status'] == 'completed'
        assert store.get('artifact', view['artifact_id'])['revision'] == 1
        assert store.get('artifact_proposal', view['proposal_id'])['status'] == 'rejected'
    finally:
        await runtime.stop()
