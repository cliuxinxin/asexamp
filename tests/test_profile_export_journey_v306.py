"""One HTTP journey through template approval, generation, process artifacts and Excel."""
import copy
import io

from openpyxl import load_workbook
from test_native_journey_v300 import native_journey


def test_template_preview_approval_and_all_process_artifacts_export(native_journey):
    j = native_journey
    before = j.client.get('/api/projects/' + j.project['id'] + '/profiles').json()[0]
    columns = copy.deepcopy(before['config']['excel_columns'])
    next(column for column in columns if column['field'] == 'title')['header'] = '演示用例名'
    original = j.gateway.generate_native

    async def model(task, context, schema, instruction):
        if task == 'learn_template':
            j.gateway.generations.append((task, copy.deepcopy(context)))
            return {'summary': '将用例标题列改名为演示用例名，并建议调整工作表名称。',
                    'template_kinds': ['cases'],
                    'config': {'excel_columns': columns, 'sheet_name': 'Template Demo'}}
        return await original(task, context, schema, instruction)

    j.gateway.generate_native = model
    uploaded = j.client.post('/api/chats/' + j.chat['id'] + '/sources',
        files={'file': ('case-template.csv', '用例编号,演示用例名\nTC-1,格式示例\n'.encode(), 'text/csv')},
        data={'role': 'example'})
    assert uploaded.status_code == 200
    j.turn('学习这份用例模板，先让我查看 Profile 更改。', 'learn_template_tool',
           {'source_ids': [uploaded.json()['id']], 'kind': 'cases'}, status='needs_confirmation')
    pending = j.snapshot()['conversation_prompt']
    path = '/api/chats/' + j.chat['id'] + '/profile-change'
    response = j.client.get(path, params={'prompt_id': pending['id']})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert {'excel_columns', 'sheet_name'} <= {row['key'] for row in preview['changes']}
    assert j.client.get('/api/projects/' + j.project['id'] + '/profiles').json()[0] == before
    calls = len(j.gateway.generations)
    applied = j.client.post(path + '/apply', json={'prompt_id': pending['id'],
        'expected_version': preview['expected_version'], 'selected_keys': ['excel_columns']})
    assert applied.status_code == 200, applied.text
    after = applied.json()['profile']
    assert after['version'] == before['version'] + 1
    assert after['config']['sheet_name'] == before['config']['sheet_name']
    assert next(c for c in after['config']['excel_columns'] if c['field'] == 'title')['header'] == '演示用例名'
    assert len(j.gateway.generations) == calls
    assert j.snapshot()['conversation_prompt'] is None

    j.turn('按已确认的模板生成测试用例。', 'start_pipeline_tool', mode='auto')
    run = j.completed()
    result = j.client.get('/api/chats/' + j.chat['id'] + '/artifacts')
    assert result.status_code == 200, result.text
    catalog = result.json()['items']
    assert {'understood', 'scenarios_generated', 'cases_generated', 'reviewed'} <= {r['phase'] for r in catalog}
    draft = next(r for r in catalog if r['phase'] == 'cases_generated')
    reviewed = next(r for r in catalog if r['phase'] == 'reviewed')
    assert draft['artifact_id'] == reviewed['artifact_id'] == run['current_artifact_id']
    assert draft['revision'] < reviewed['revision']
    assert not draft['is_current'] and reviewed['is_current']
    for entry in (draft, reviewed):
        exported = j.client.get('/api/artifacts/' + entry['artifact_id'] + '/export',
                               params={'revision': entry['revision']})
        assert exported.status_code == 200, exported.text[:300] if exported.status_code != 200 else ''
        sheet = load_workbook(io.BytesIO(exported.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        column = rows[0].index('演示用例名')
        assert str(rows[1][column]).startswith('已评审：') == (entry['phase'] == 'reviewed')
    assert j.snapshot()['runs'][0]['status'] == 'completed'
