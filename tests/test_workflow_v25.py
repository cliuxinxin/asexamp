from fastapi.testclient import TestClient
from tcg.main import create_app
from test_backend_api import Model, setup_chat, start, until

class FlowModel(Model):
    async def generate(self, task, context):
        if task=='summarize':
            self.calls.append((task,context))
            return {'summary':'覆盖登录正常路径；请按实际环境执行并复核。'}
        if task=='dialogue':
            self.calls.append((task,context))
            return {'answer':'人工模式在理解与场景确认节点等待。','refs':[]}
        if task=='route':
            self.calls.append((task,context))
            return {'intent':'query' if '怎么' in context['request']['content'] else 'generate_case'}
        result = await super().generate(task, context)
        if task == 'analyze_requirement':
            result['report']['summary']='理解登录规则'
            result['report']['strategy']={'depth':'standard','rationale':'覆盖正常和错误凭证'}
            result['report']['diagrams']=[{'title':'业务流程','mermaid':'flowchart TD\n A[登录] --> B[验证凭证]'}]
            if context.get('clarification'):result['report']['questions']=[]
        if task == 'generate_scenarios':
            for item in result['items']:item['requirement_ids']=[i['id'] for i in context['analysis']]
        return result

def test_hitp_gates_and_latest_scenario(tmp_path):
    model=FlowModel(questions=['允许哪些角色登录？'])
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client)
        run=until(client,start(client,chat,experience='reliable',mode='hitp'))
        assert run['status']=='waiting' and run['interrupt']['type']=='clarification',run
        assert not any(t=='generate_cases' for t,_ in model.calls)
        client.post('/api/runs/'+run['id']+'/resume',json={'answer':'注册用户'})
        run=until(client,run)
        assert run['interrupt']['type']=='strategy_review',run
        client.post('/api/runs/'+run['id']+'/resume',json={'approved':True})
        run=until(client,run)
        assert run['interrupt']['type']=='scenario_review',run
        artifact=client.get('/api/artifacts/'+run['interrupt']['artifact_id']).json()
        artifact['items'][0]['title']='人工修改后的场景'
        assert client.put('/api/artifacts/'+artifact['id'],json={'expected_revision':artifact['revision'],'items':artifact['items']}).status_code==200
        client.post('/api/runs/'+run['id']+'/resume',json={'approved':True})
        run=until(client,run)
        assert run['status']=='completed',run
        cases=[c for t,c in model.calls if t=='generate_cases']
        assert cases[0]['scenarios'][0]['title']=='人工修改后的场景'
        assert sum(t=='review_cases' for t,_ in model.calls)==1

def test_auto_keeps_stages_and_no_pauses(tmp_path):
    model=FlowModel(questions=['未说明错误提示文案'])
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client, 'Valid credentials log in.\n\n' * 100)
        run=until(client,start(client,chat,experience='reliable',mode='auto'))
        assert run['status']=='completed',run
        assert [t for t,_ in model.calls]==['analyze_requirement','generate_scenarios','generate_cases','review_cases','summarize']
        visible=client.get('/api/chats/'+chat['id']).json()['messages']
        assert len([m for m in visible if m['role']=='assistant'])==1
        assert client.get('/api/artifacts/'+run['artifact_ids'][0]+'/export').status_code==200


def test_profile_override_and_excel_columns(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    model=FlowModel()
    with TestClient(create_app(tmp_path,model)) as client:
        project,chat,_=setup_chat(client)
        profiles=client.get('/api/projects/'+project['id']+'/profiles').json()
        config={'scenario_level':'deep','case_level':'quick','sheet_name':'验收模板','excel_columns':[{'field':'title','header':'用例名称'},{'field':'expected','header':'预期结果'}]}
        run=until(client,start(client,chat,experience='reliable',mode='auto',profile_id=profiles[0]['id'],profile_override=config))
        assert run['status']=='completed',run
        exported=client.get('/api/artifacts/'+run['artifact_ids'][0]+'/export')
        assert exported.status_code==200
        sheet=load_workbook(BytesIO(exported.content)).active
        assert sheet.title=='验收模板'
        assert list(next(sheet.values))==['用例名称','预期结果']
        assert sheet.cell(2,1).value=='Reviewed login case'
        assert client.get('/api/projects/'+project['id']+'/profiles').json()==profiles


def test_explicit_review_preserves_target_and_records_changes(tmp_path):
    model=FlowModel()
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client)
        generated=until(client,start(client,chat,experience='reliable'))
        aid=generated['artifact_ids'][0]
        before=client.get('/api/artifacts/'+aid).json()
        reviewed=until(client,start(client,chat,experience='reliable',intent='review_case',artifact_id=aid))
        assert reviewed['status']=='completed',reviewed
        after=client.get('/api/artifacts/'+aid).json()
        assert after['revision']>before['revision']
        assert after['report']['review_reports'][0]['changes']['update']==1


def test_clarification_retry_reuses_answer(tmp_path):
    class RecoverModel(FlowModel):
        broken=True
        async def generate(self,task,context):
            result=await super().generate(task,context)
            if task=='analyze_requirement' and context.get('clarification') and self.broken:
                result['report']['diagrams']=[]
            return result
    model=RecoverModel(questions=['允许哪些角色？'])
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client)
        run=until(client,start(client,chat,experience='reliable',mode='hitp'))
        client.post('/api/runs/'+run['id']+'/resume',json={'answer':'注册用户'})
        run=until(client,run)
        assert run['status']=='failed',run
        model.broken=False
        assert client.post('/api/runs/'+run['id']+'/retry',json={}).status_code==200
        run=until(client,run)
        assert run['status']=='waiting' and run['interrupt']['type']=='strategy_review',run
        sources=client.get('/api/chats/'+chat['id']).json()['sources']
        assert len([s for s in sources if s['role']=='clarification'])==1
        last=[c for t,c in model.calls if t=='analyze_requirement'][-1]
        assert 'validation_repair' in last


def test_example_source_confirmation_then_generate(tmp_path):
    model=FlowModel()
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,source=setup_chat(client)
        client.put('/api/sources/'+source['id']+'/role',json={'role':'example'})
        run=until(client,start(client,chat,experience='reliable'))
        assert run['status']=='waiting' and run['interrupt']['type']=='source_review',run
        assert not model.calls
        messages=client.get('/api/chats/'+chat['id']).json()['messages']
        assert not any('共 1 条用例' in m['content'] for m in messages)
        response=client.post('/api/runs/'+run['id']+'/resume',json={'source_ids':[source['id']]})
        assert response.status_code==200,response.text
        run=until(client,run)
        assert run['status']=='completed',run
        assert client.get('/api/artifacts/'+run['artifact_ids'][0]).json()['type']=='cases'


def test_chat_without_documents_and_paused_question(tmp_path):
    model=FlowModel(questions=['允许哪些角色？'])
    with TestClient(create_app(tmp_path,model)) as client:
        project=client.get('/api/projects').json()[0]
        chat=client.post('/api/projects/'+project['id']+'/chats',json={}).json()
        run=until(client,start(client,chat,experience='reliable',intent='auto',content='怎么使用人工模式？'))
        assert run['status']=='completed',run
        artifact=client.get('/api/artifacts/'+run['artifact_ids'][0]).json()
        assert artifact['type']=='answer'
        assert not any(t=='analyze_requirement' for t,c in model.calls)
        client.post('/api/chats/'+chat['id']+'/sources/text',json={'name':'需求','text':'有效账号允许登录','role':'primary'})
        run=until(client,start(client,chat,experience='reliable',mode='hitp'))
        assert run['status']=='waiting'
        before=run['interrupt']
        response=client.post('/api/runs/'+run['id']+'/dialogue',json={'content':'为什么需要确认？'})
        assert response.status_code==200,response.text
        after=client.get('/api/runs/'+run['id']).json()
        assert after['status']=='waiting' and after['interrupt']==before


def test_continue_from_scenarios_without_reanalysis(tmp_path):
    model=FlowModel()
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client)
        run=until(client,start(client,chat,experience='reliable',intent='generate_scenario'))
        assert run['status']=='completed',run
        aid=run['artifact_ids'][0]
        model.calls.clear()
        run=until(client,start(client,chat,experience='reliable',artifact_id=aid))
        assert run['status']=='completed',run
        assert [t for t,c in model.calls]==['generate_cases','review_cases','summarize']


def test_template_proposal_and_modification(tmp_path):
    model=FlowModel()
    with TestClient(create_app(tmp_path,model)) as client:
        project,chat,_=setup_chat(client)
        profiles=client.get('/api/projects/'+project['id']+'/profiles').json()
        learned=until(client,start(client,chat,experience='reliable',intent='learn_template'))
        assert learned['status']=='completed',learned
        assert client.get('/api/artifacts/'+learned['artifact_ids'][0]).json()['type']=='proposal'
        assert client.get('/api/projects/'+project['id']+'/profiles').json()==profiles
        generated=until(client,start(client,chat,experience='reliable'))
        aid=generated['artifact_ids'][0]
        modified=until(client,start(client,chat,experience='reliable',intent='modify',artifact_id=aid,content='修改标题'))
        assert modified['status']=='completed',modified
        artifact=client.get('/api/artifacts/'+aid).json()
        assert artifact['items'][0]['title']=='Modified through chat'


def test_checkpoint_restart_resumes_same_confirmation(tmp_path):
    model=FlowModel(questions=['允许哪些角色？'])
    with TestClient(create_app(tmp_path,model)) as client:
        _,chat,_=setup_chat(client)
        run=until(client,start(client,chat,experience='reliable',mode='hitp'))
        assert run['interrupt']['type']=='clarification'
    with TestClient(create_app(tmp_path,model)) as client:
        assert client.get('/api/runs/'+run['id']).json()['status']=='waiting'
        response=client.post('/api/runs/'+run['id']+'/resume',json={'answer':'注册用户'})
        assert response.status_code==200
        run=until(client,run)
        assert run['interrupt']['type']=='strategy_review'
        assert sum(t=='analyze_requirement' for t,c in model.calls)==2
        sources=client.get('/api/chats/'+chat['id']).json()['sources']
        assert sum(s['role']=='clarification' for s in sources)==1
