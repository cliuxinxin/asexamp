"""Continuous conversation integration with real SQLite and registered services.

Only semantic model output is controlled. Generation/HTTP runtime is not mocked
or claimed here: the conversation begins with genuinely committed fixtures.
The full 16-step live acceptance input is fixtures/conversation_v260.json.
"""
import copy
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from tcg.conversation import ConversationController
from tcg.documents import parse_text
from tcg.schemas import MessageInput
from tcg.storage import Store, now


class ControlledModelEngine:
    """Deterministic semantics; every write still crosses real domain services."""

    def __init__(self, store):
        self.store = store
        self.decision = None
        self.calls = []
        self.hook = None

    def fits(self, task, context):
        return True

    async def invoke_model(self, task, context, run_id=None):
        self.calls.append((task, copy.deepcopy(context), run_id))
        if self.hook:
            self.hook(task, context)
        if task in ('conversation_interpret', 'conversation_turn'):
            if self.decision is None:
                raise AssertionError('No semantic decision supplied')
            return copy.deepcopy(self.decision)
        if task == 'artifact_estimate':
            return {'scenarios': [
                {'scenario_id': row['id'], 'min_count': 1, 'max_count': 3,
                 'rationale': '当前已提交场景的异常和边界路径',
                 'assumptions': ['仅设计估算，未生成或执行用例']}
                for row in context['scenarios']]}
        if task == 'learn_template':
            reference = context['format_references'][0]['text']
            if 'SCENARIO-TEMPLATE' in reference:
                return {'template_kinds': ['scenarios'], 'config': {
                    'scenario_sheet_name': '场景模板工作表', 'scenario_excel_columns': [
                        {'field': 'id', 'header': '场景号'}, {'field': 'title', 'header': '业务场景'}]},
                    'summary': '已识别场景模板列。'}
            if 'CASE-TEMPLATE' in reference:
                return {'template_kinds': ['cases'], 'config': {
                    'sheet_name': '用例模板工作表', 'excel_columns': [
                        {'field': 'id', 'header': '用例号'}, {'field': 'title', 'header': '业务用例'},
                        {'field': 'human_note', 'header': '人工备注', 'value_source': 'manual'}]},
                    'summary': '已识别用例模板和人工字段。'}
            raise AssertionError('Unknown template fixture')
        if task == 'project_source_impact':
            return {'requirement_ids': [r['id'] for r in context['requirements'] if r['id'] == 'R-2'],
                    'summary': '新增强制下线规则只影响账号锁定需求及其场景。',
                    'global_impact': False, 'uncertain': False,
                    'refs': [row['id'] for row in context['new_evidence']]}
        if task in ('artifact_explain', 'artifact_analyze', 'conversation_answer'):
            rows = context.get('artifact', {}).get('items', [])
            refs = [item['id'] for item in context.get('evidence', []) if item.get('role') != 'example']
            return {'answer': '所选条目：' + '、'.join(row['id'] for row in rows), 'refs': refs}
        if task in ('artifact_modify', 'artifact_sync', 'artifact_sync_scenarios'):
            rows = context['artifact']['items']
            operations = []
            for row in rows:
                item = {'id': row['id'], 'title': row['title'] + '（说明已更新）'}
                if '强制下线' in context.get('instruction', ''):
                    item['description'] = '账号锁定会注销其他活动会话；解除后失败计数归零。'
                    item['refs'] = list(dict.fromkeys(row.get('refs', []) + [
                        e['id'] for e in context.get('evidence', []) if e.get('role') == 'change']))
                if context['artifact']['type'] == 'cases':
                    item['steps'] = [{'action': '逐项执行并核对输入', 'expected': '实际行为满足对应需求'}]
                    item['human_note'] = 'MODEL-MUST-NOT-OVERWRITE-MANUAL-NOTE'
                operations.append({'op': 'update', 'id': row['id'], 'item': item})
            return {'operations': operations, 'summary': '根据所选条目生成明确修改。'}
        if task in ('artifact_review', 'artifact_review_cases', 'artifact_review_readonly', 'review_cases'):
            return {'operations': [], 'report': {'summary': '仅评审当前用例，没有修改用例。', 'issues': []}}
        raise AssertionError(f'Unexpected model task: {task}; keys={list(context)}')


class ContinuousConversationTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.store = Store(self.directory.name)
        self.project = self.store.list('project')[0]
        self.profile = self.store.list('profile')[0]
        columns = [{'field': 'id', 'header': '用例编号'}, {'field': 'title', 'header': '用例名称'},
                   {'field': 'steps', 'header': '操作步骤'}, {'field': 'expected', 'header': '预期结果'},
                   {'field': 'human_note', 'header': '人工备注', 'value_source': 'manual'}]
        self.profile = self.store.update_profile(self.profile['id'], self.profile['name'],
            {**self.profile['config'], 'excel_columns': columns, 'sheet_name': '最终用例',
             'scenario_excel_columns': [{'field': 'id', 'header': '场景编号'},
                                        {'field': 'title', 'header': '场景名称'}],
             'scenario_sheet_name': '测试场景'}, self.profile['version'])
        self.chat = self.store.create_chat(self.project['id'], 'continuous A')
        self.other_chat = self.store.create_chat(self.project['id'], 'shared B')
        self.fixture = json.loads((ROOT / 'tests/fixtures/conversation_v260.json').read_text())
        content, chunks = parse_text(self.fixture['requirement'])
        self.source = self.store.add_source(self.chat['id'], '登录需求', 'primary', content, chunks)
        self.ref = self.store.evidence([self.source['id']])[0]['id']
        request = MessageInput(content='seed committed inputs', experience='reliable',
            intent='generate_case', profile_id=self.profile['id']).model_dump()
        _, run = self.store.create_run(self.chat['id'], request)
        self.run_id = run['id']
        rules = [{'id': f'R-{i}', 'title': title, 'description': title, 'refs': [self.ref]}
                 for i, title in enumerate(('正确密码登录', '账号锁定', '主动退出'), 1)]
        self.analysis = self.store.artifact(run['id'], 'seed-analysis', 'analysis', '需求理解', rules, {})
        scenes = [{'id': f'S-{i}', 'title': row['title'], 'description': row['description'],
                   'priority': 'P1', 'refs': [self.ref], 'requirement_ids': [row['id']]}
                  for i, row in enumerate(rules, 1)]
        self.scenarios = self.store.artifact(run['id'], 'seed-scenes', 'scenarios', '测试场景', scenes, {})
        self.scenarios['report'] = {'lineage': {'analysis_artifact_id': self.analysis['id'], 'analysis_revision': 1}}
        cases = [{'id': f'C-{i}', 'title': row['title'], 'scenario_id': row['id'], 'type': 'Business',
                  'priority': 'P1', 'preconditions': '使用已注册的测试账号',
                  'steps': [{'action': '执行对应用户操作', 'expected': '观察到需求约定的结果'}],
                  'refs': [self.ref], 'human_note': self.fixture['human_note']}
                 for i, row in enumerate(scenes, 1)]
        self.cases = self.store.artifact(run['id'], 'seed-cases', 'cases', '测试用例', cases, {})
        self.cases['report'] = {'lineage': {'scenario_artifact_id': self.scenarios['id'], 'scenario_revision': 1}}
        for artifact in (self.analysis, self.scenarios, self.cases):
            artifact['_visible'] = True
            self.store.put('artifact', artifact)
            from tcg.storage import dump
            self.store.db.execute('UPDATE revisions SET payload=? WHERE artifact_id=? AND revision=1',
                                 (dump(artifact), artifact['id']))
        self.store.update_run(run['id'], status='completed', stage='completed',
            artifact_ids=[self.cases['id']])
        self.engine = ControlledModelEngine(self.store)
        self.controller = ConversationController(self.store, self.engine)
        self.counter = 0

    def tearDown(self):
        self.store.close()
        self.directory.cleanup()

    async def asyncTearDown(self):
        await self.controller.close()

    async def ask(self, content, actions, **hints):
        self.counter += 1
        self.engine.decision = {'actions': copy.deepcopy(actions)}
        result = await self.controller.submit(self.chat['id'], {
            'client_message_id': f'continuous-{self.counter}', 'content': content, **hints})
        self.assertNotEqual(result['status'], 'failed', result)
        self.assertEqual(self.controller.get(self.chat['id'], result['id']), result)
        return result

    def artifact_args(self, artifact, **extra):
        current = self.store.get('artifact', artifact['id'])
        return {'artifact_id': artifact['id'], 'expected_revision': current['revision'], **extra}

    def action(self, name, **arguments):
        return {'name': name, 'arguments': arguments}

    def part(self, response, kind):
        values = [part for part in response['parts'] if part['type'] == kind]
        self.assertTrue(values, response)
        return values[0]

    async def test_seeded_continuous_read_and_versioned_write_chain(self):
        before = self.store.get('artifact', self.cases['id'])
        result = await self.ask('展开第三条用例的步骤和预期', [self.action('artifact.read',
            **self.artifact_args(self.cases, selected_ids=['C-3'], detail='steps'))],
            artifact_id=self.cases['id'], artifact_revision=1, selected_ids=['C-3'])
        details = [part for part in result['parts'] if part['type'] == 'case_details']
        self.assertTrue(details, result)
        self.assertEqual(details[0]['items'][0]['id'], 'C-3')
        self.assertEqual(details[0]['items'][0]['steps'], before['items'][2]['steps'])
        self.assertEqual(self.store.get('artifact', self.cases['id']), before)

        scene_before = self.store.get('artifact', self.scenarios['id'])
        estimate = await self.ask('只估第二、第三个场景要多少条用例，先不生成', [
            self.action('artifact.estimate', artifact_id=self.scenarios['id'], ordinals=[2, 3],
                        instruction='只估当前第二、第三个场景，不生成')],
            artifact_id=self.scenarios['id'], artifact_revision=1, view_order=['S-3', 'S-2', 'S-1'])
        data = self.part(estimate, 'estimate')['data']
        self.assertEqual({row['scenario_id'] for row in data['scenarios']}, {'S-2', 'S-1'})
        self.assertEqual((data['min_count'], data['max_count']), (2, 6))
        followup = await self.ask('那只算异常的呢？', [self.action('artifact.estimate',
            scope='inherit_previous', instruction='仅估算异常路径')])
        self.assertEqual({row['scenario_id'] for row in self.part(followup, 'estimate')['data']['scenarios']},
                         {'S-2', 'S-1'})
        runs_before = copy.deepcopy(self.store.runs(self.chat['id']))
        await self.ask('可以', [self.action('workflow.continue', confirmation='explicit')])
        self.assertEqual(self.store.runs(self.chat['id']), runs_before)
        self.assertEqual(self.store.get('artifact', self.scenarios['id']), scene_before)

        explained = await self.ask('解释第三条用例，保留原文', [self.action('artifact.analyze',
            **self.artifact_args(self.cases, selected_ids=['C-3'], instruction='解释第三条用例'))])
        self.assertIn('C-3', self.part(explained, 'answer')['text'])
        preview = await self.ask('把第三条的步骤写清楚，保留人工备注，先不保存', [
            self.action('artifact.preview', **self.artifact_args(self.cases,
                selected_ids=['C-3'], instruction='完善第三条步骤，保留人工备注'))])
        proposal = self.part(preview, 'diff')['proposal_id']
        self.assertEqual(self.store.get('artifact', self.cases['id']), before)
        applied = await self.ask('应用这项修改', [self.action('artifact.apply', proposal_id=proposal,
                                                           artifact_id=self.cases['id'])])
        case_after = self.store.get('artifact', self.cases['id'])
        self.assertEqual(case_after['revision'], 2)
        self.assertEqual(case_after['items'][:2], before['items'][:2])
        self.assertEqual(case_after['items'][2]['human_note'], self.fixture['human_note'])
        self.assertNotEqual(case_after['items'][2]['steps'], before['items'][2]['steps'])
        self.assertEqual(case_after['items'][2]['refs'], before['items'][2]['refs'])

        linked = await self.ask('把刚才用例对应的场景写清楚，并同步关联用例；其他场景别动，先看差异', [
            self.action('artifact.preview', **self.artifact_args(self.scenarios,
                selected_ids=['S-3'], instruction='完善第三条场景并同步关联用例', sync_related=True))])
        linked_diff = self.part(linked, 'diff')
        self.assertEqual({change['artifact_id'] for change in linked_diff['changes']},
                         {self.scenarios['id'], self.cases['id']})
        self.assertEqual(self.store.get('artifact', self.scenarios['id']), scene_before)
        self.assertEqual(self.store.get('artifact', self.cases['id']), case_after)
        await self.ask('应用刚才的联动修改', [self.action('artifact.apply',
            artifact_id=self.scenarios['id'], proposal_id=linked_diff['proposal_id'])])
        scene_after = self.store.get('artifact', self.scenarios['id'])
        synced_cases = self.store.get('artifact', self.cases['id'])
        self.assertEqual(scene_after['revision'], 2)
        self.assertEqual(scene_after['items'][:2], scene_before['items'][:2])
        self.assertEqual(synced_cases['items'][:2], case_after['items'][:2])
        self.assertEqual(synced_cases['items'][2]['human_note'], self.fixture['human_note'])
        sync_context = [context for task, context, _ in self.engine.calls if task == 'artifact_sync'][-1]
        self.assertEqual(sync_context['scenarios'], [scene_after['items'][2]])

        coverage = await self.ask('展示需求、场景和用例的覆盖，哪些缺用例、哪些没同步？', [
            self.action('artifact.coverage', **self.artifact_args(self.cases))])
        coverage_data = self.part(coverage, 'coverage')['data']
        self.assertIn('coverage', coverage_data)
        self.assertEqual(coverage_data['coverage']['totals']['cases'], 3)
        review = await self.ask('只评审这些用例，不修改', [self.action('artifact.review_cases',
            **self.artifact_args(self.cases, instruction='只评审，不修改'))])
        self.assertTrue(review['parts'])
        self.assertEqual(self.store.get('artifact', self.cases['id']), synced_cases)

        await self.ask('把当前第一条用例写法保存为本项目 Profile 的样例', [self.action('project.pin_samples',
            **self.artifact_args(self.cases, profile_id=self.profile['id'],
                expected_version=self.profile['version'], selected_ids=['C-1']))])
        profile = self.store.get('profile', self.profile['id'])
        self.assertEqual(profile['version'], self.profile['version'] + 1)
        sample = profile['config']['sample_cases'][0]
        self.assertEqual(sample['steps'], synced_cases['items'][0]['steps'])
        self.assertNotIn('refs', sample)
        self.assertNotIn('scenario_id', sample)
        _, second_run = self.store.create_run(self.other_chat['id'], MessageInput(
            content='Use saved format', experience='reliable', intent='generate_case',
            profile_id=profile['id']).model_dump())
        self.assertEqual(second_run['_profile']['sample_cases'], profile['config']['sample_cases'])
        self.assertEqual(second_run['_source_ids'], [])

        summary = await self.ask('用产品经理能理解的方式总结这些用例，再列出相关编号；不要修改', [
            self.action('artifact.analyze', **self.artifact_args(self.cases, instruction='总结当前用例并列出编号'))])
        answer = self.part(summary, 'answer')['text']
        for row in synced_cases['items']:
            self.assertIn(row['id'], answer)
        self.assertEqual(self.store.get('artifact', self.cases['id']), synced_cases)

        exported = await self.ask('把最新场景和最终用例分别导出 Excel', [self.action('artifact.export',
            artifact_ids=[self.scenarios['id'], self.cases['id']], profile_id=profile['id'])])
        files = self.part(exported, 'files')['files']
        self.assertEqual(len(files), 2)
        self.assertEqual(len({item['url'] for item in files}), 2)
        from openpyxl import load_workbook
        from tcg.conversation_project import frozen_export
        frozen_files = {}
        for item in files:
            export_id = item['url'].rsplit('/', 1)[-1]
            record, payload = frozen_export(self.store, export_id)
            self.assertTrue(payload.startswith(b'PK'))
            sheet = load_workbook(io.BytesIO(payload)).active
            if item['artifact_id'] == self.scenarios['id']:
                self.assertEqual(sheet.title, '测试场景')
                self.assertEqual(list(next(sheet.values)), ['场景编号', '场景名称'])
                self.assertEqual(record['revision'], scene_after['revision'])
            else:
                self.assertEqual(sheet.title, '最终用例')
                self.assertEqual(list(next(sheet.values)), ['用例编号', '用例名称', '操作步骤', '预期结果', '人工备注'])
                self.assertEqual(record['revision'], synced_cases['revision'])
            self.assertEqual(sheet.max_row, 4)
            frozen_files[export_id] = payload

        last_preview = await self.ask('第三条再写简洁些，先预览', [self.action('artifact.preview',
            **self.artifact_args(self.cases, selected_ids=['C-3'], instruction='简化第三条标题'))])
        last_proposal = self.part(last_preview, 'diff')['proposal_id']
        await self.ask('取消这个修改，其他任务保留', [self.action('artifact.discard',
            artifact_id=self.cases['id'], proposal_id=last_proposal)])
        self.assertEqual(self.store.get('artifact', self.cases['id']), synced_cases)
        self.assertEqual(self.store.run(self.run_id)['status'], 'completed')
        self.assertEqual(self.counter, 16)
        for export_id, payload in frozen_files.items():
            self.assertEqual(frozen_export(self.store, export_id)[1], payload)

    async def test_chat_adopt_card_edit_reopen_chat_save_share_has_one_authoritative_answer(self):
        from tcg.clarification import get_draft, update_draft
        from tcg.project_context import shared_sources

        _, waiting = self.store.create_run(self.chat['id'], MessageInput(
            content='clarification fixture', experience='reliable', intent='generate_scenario',
            mode='hitp').model_dump())
        question = self.fixture['clarification_question']
        self.store.update_run(waiting['id'], status='waiting', stage='clarification',
            _interrupt_id='integration-real-draft-gate', stop_after='scenarios',
            interrupt={'type': 'clarification', 'questions': [question], 'question_suggestions': [{
                'question': question, 'answer': self.fixture['suggested_answer'],
                'basis': '待确认的测试设计假设', 'confidence': 'assumption', 'refs': []}]})
        draft = get_draft(self.store, waiting['id'])
        adopted = await self.ask('采用全部建议，但先别继续', [self.action('clarification.adopt',
            run_id=waiting['id'], adopt_all=True, expected_revision=draft['revision'])])
        self.assertTrue(any(part['type'] == 'clarification_draft' for part in adopted['parts']))
        draft = get_draft(self.store, waiting['id'])
        self.assertTrue(draft['questions'][0]['adopted'])
        self.assertFalse(draft['submitted'])
        self.assertEqual(shared_sources(self.store, self.project['id']), [])

        # This is the same service invoked by the card PATCH endpoint.
        edited = update_draft(self.store, waiting['id'], {'expected_revision': draft['revision'],
            'answers': {draft['questions'][0]['id']: self.fixture['confirmed_answer']}})
        await self.controller.close()
        self.store.close()
        self.store = Store(self.directory.name)
        self.engine.store = self.store
        self.controller = ConversationController(self.store, self.engine)
        restored = get_draft(self.store, waiting['id'])
        self.assertEqual(restored, edited)
        self.assertTrue(restored['questions'][0]['adopted'])

        saved = await self.ask('按卡片里的十分钟答案保存到项目，先不继续', [
            self.action('clarification.save', run_id=waiting['id'], expected_revision=restored['revision']),
            self.action('clarification.share', run_id=waiting['id'])])
        draft = get_draft(self.store, waiting['id'])
        self.assertTrue(draft['submitted'])
        self.assertTrue(draft['shared'])
        self.assertEqual(draft['questions'][0]['answer'], self.fixture['confirmed_answer'])
        shared = shared_sources(self.store, self.project['id'])
        self.assertEqual(len(shared), 1)
        self.assertIn(self.fixture['confirmed_answer'], shared[0]['_text'])
        self.assertNotIn(self.fixture['suggested_answer'], shared[0]['_text'])
        run = self.store.run(waiting['id'])
        self.assertEqual(run['status'], 'waiting')
        self.assertEqual(run['_interrupt_id'], 'integration-real-draft-gate')
        self.assertEqual(run['stop_after'], 'scenarios')
        self.assertEqual(self.controller.get(self.chat['id'], saved['id']), saved)

        # The persisted source enters B's real Store-created run exactly once.
        _, second = self.store.create_run(self.other_chat['id'], MessageInput(
            content='reuse confirmed knowledge', experience='reliable', intent='generate_case').model_dump())
        self.assertIn(draft['source_id'], second['_source_ids'])
        shared_refs = [row['id'] for row in self.store.evidence(second['_source_ids'])]
        self.assertTrue(shared_refs)
        self.assertTrue(all(ref.startswith(draft['source_id'] + '#') for ref in shared_refs))

    async def test_actual_write_receipt_survives_reopen_and_old_preview_cannot_overwrite(self):
        body = {'client_message_id': 'persisted-actual-write', 'content': '直接写清第三条用例',
                'artifact_id': self.cases['id'], 'artifact_revision': 1}
        self.engine.decision = {'actions': [self.action('artifact.revise',
            **self.artifact_args(self.cases, selected_ids=['C-3'], instruction='完善第三条'))]}
        first = await self.controller.submit(self.chat['id'], body)
        self.assertEqual(first['status'], 'succeeded', first)
        self.assertEqual(self.store.get('artifact', self.cases['id'])['revision'], 2)
        calls = len(self.engine.calls)
        await self.controller.close()
        self.store.close()
        self.store = Store(self.directory.name)
        self.engine.store = self.store
        self.controller = ConversationController(self.store, self.engine)
        retried = await self.controller.submit(self.chat['id'], body)
        self.assertEqual(retried, first)
        self.assertEqual(len(self.engine.calls), calls)
        self.assertEqual(self.store.get('artifact', self.cases['id'])['revision'], 2)

        preview = await self.ask('再简洁些，先预览', [self.action('artifact.preview',
            **self.artifact_args(self.cases, selected_ids=['C-3'], instruction='简化第三条'))])
        proposal_id = self.part(preview, 'diff')['proposal_id']
        current = self.store.get('artifact', self.cases['id'])
        manually_edited = copy.deepcopy(current['items'])
        manually_edited[2]['human_note'] = '用户在预览后补充了人工备注。'
        latest = self.store.revise_artifact(current['id'], current['revision'], manually_edited)
        self.engine.decision = {'actions': [self.action('artifact.apply',
            artifact_id=current['id'], proposal_id=proposal_id)]}
        conflict = await self.controller.submit(self.chat['id'], {
            'client_message_id': 'apply-stale-preview', 'content': '应用旧预览'})
        self.assertNotEqual(conflict['status'], 'succeeded', conflict)
        self.assertEqual(self.store.get('artifact', current['id']), latest)

    async def test_read_snapshot_finishes_after_main_workflow_advances_without_edit_lock(self):
        _, running = self.store.create_run(self.chat['id'], MessageInput(
            content='active run fixture', experience='reliable', intent='generate_case').model_dump())
        self.store.update_run(running['id'], status='running', stage='case_generation')
        before = self.store.get('artifact', self.cases['id'])
        observed = []

        def advance(task, context):
            if task == 'artifact_explain':
                observed.append(self.store.run(running['id']).get('_edit_token'))
                self.assertFalse(getattr(self.store, '_workspace_action_tokens', {}))
                self.store.update_run(running['id'], status='waiting', stage='case_review',
                    _interrupt_id='new-boundary', interrupt={'type': 'scenario_review',
                                                            'artifact_id': self.scenarios['id']})

        self.engine.hook = advance
        result = await self.ask('生成过程中先解释已有用例，不改变主流程', [self.action('artifact.analyze',
            **self.artifact_args(self.cases, selected_ids=['C-3'], instruction='解释已有第三条'))])
        self.assertEqual(result['status'], 'succeeded', result)
        self.assertEqual(observed, [None])
        self.assertIn('C-3', self.part(result, 'answer')['text'])
        self.assertEqual(self.store.get('artifact', self.cases['id']), before)
        self.assertEqual(self.store.run(running['id'])['_interrupt_id'], 'new-boundary')

    async def test_template_learning_and_application_keep_distinct_xlsx_columns_and_manual_policy(self):
        template_ids = []
        for marker in ('SCENARIO-TEMPLATE', 'CASE-TEMPLATE'):
            text, chunks = parse_text(marker)
            template_ids.append(self.store.add_source(self.chat['id'], marker, 'example', text, chunks)['id'])
        before_runs = copy.deepcopy(self.store.runs(self.chat['id']))
        learned = await self.ask('分别学习两份模板，并应用到当前 Profile', [
            self.action('project.learn_template', source_ids=template_ids, kind='both', apply=True,
                        profile_id=self.profile['id'], expected_version=self.profile['version'])])
        self.assertEqual(learned['status'], 'succeeded', learned)
        profile = self.store.get('profile', self.profile['id'])
        self.assertEqual(profile['version'], self.profile['version'] + 1)
        self.assertEqual(profile['config']['scenario_sheet_name'], '场景模板工作表')
        self.assertEqual(profile['config']['sheet_name'], '用例模板工作表')
        self.assertEqual(profile['config']['excel_columns'][-1]['value_source'], 'manual')
        self.assertEqual(self.store.runs(self.chat['id']), before_runs)
        contexts = [context for task, context, _ in self.engine.calls if task == 'learn_template']
        self.assertEqual(len(contexts), 2)
        self.assertTrue(all(context['evidence'] == [] for context in contexts))
        self.assertTrue(all(all(e['role'] == 'example' for e in context['format_references']) for context in contexts))

        await self.ask('只重新应用场景模板，用例模板和人工备注策略保留', [
            self.action('project.learn_template', source_ids=template_ids[:1], kind='scenario', apply=True,
                        profile_id=profile['id'], expected_version=profile['version'])])
        second = self.store.get('profile', profile['id'])
        self.assertEqual(second['config']['excel_columns'], profile['config']['excel_columns'])
        self.assertEqual(second['config']['sheet_name'], profile['config']['sheet_name'])
        self.assertEqual(second['version'], profile['version'] + 1)
        exported = await self.ask('按刚学的两份模板分别导出最新场景和用例', [self.action('artifact.export',
            artifact_ids=[self.scenarios['id'], self.cases['id']], profile_id=profile['id'])])
        from openpyxl import load_workbook
        from tcg.conversation_project import frozen_export
        sheets = {}
        for file in self.part(exported, 'files')['files']:
            _, payload = frozen_export(self.store, file['url'].rsplit('/', 1)[-1])
            sheet = load_workbook(io.BytesIO(payload)).active
            sheets[sheet.title] = list(next(sheet.values))
        self.assertEqual(sheets, {'场景模板工作表': ['场景号', '业务场景'],
                                  '用例模板工作表': ['用例号', '业务用例', '人工备注']})
        self.assertEqual(self.store.runs(self.chat['id']), before_runs)

    async def test_source_change_updates_only_affected_analysis_and_scenarios_preserving_stop(self):
        text, chunks = parse_text(self.fixture['change'])
        change = self.store.add_source(self.chat['id'], '强制下线说明', 'change', text, chunks)
        _, waiting = self.store.create_run(self.chat['id'], MessageInput(
            content='waiting at scenarios fixture', experience='reliable', intent='generate_case',
            mode='hitp', source_ids=[self.source['id']]).model_dump())
        self.store.update_run(waiting['id'], status='waiting', stage='scenario_review',
            stop_after='scenarios', _interrupt_id='scenario-confirmation',
            interrupt={'type': 'scenario_review', 'artifact_id': self.scenarios['id']})
        analysis_before = self.store.get('artifact', self.analysis['id'])
        scenarios_before = self.store.get('artifact', self.scenarios['id'])
        cases_before = self.store.get('artifact', self.cases['id'])
        result = await self.ask('结合强制下线说明更新需求理解和场景，仍然先别生成用例', [
            self.action('project.update_from_sources', **self.artifact_args(self.analysis,
                source_ids=[change['id']], targets=['analysis', 'scenarios'],
                instruction='结合强制下线说明定向更新账号锁定要求和相关场景'))])
        self.assertEqual(result['status'], 'succeeded', result)
        analysis = self.store.get('artifact', self.analysis['id'])
        scenarios = self.store.get('artifact', self.scenarios['id'])
        for old, new in ((analysis_before, analysis), (scenarios_before, scenarios)):
            self.assertEqual(new['revision'], old['revision'] + 1)
            self.assertEqual(new['items'][0], old['items'][0])
            self.assertEqual(new['items'][2], old['items'][2])
            self.assertIn('注销其他活动会话', new['items'][1]['description'])
            self.assertTrue(any(ref.startswith(change['id'] + '#') for ref in new['items'][1]['refs']))
            self.assertIn(change['id'], new['_source_ids'])
        self.assertEqual(self.store.get('artifact', self.cases['id']), cases_before)
        run = self.store.run(waiting['id'])
        self.assertEqual(run['status'], 'waiting')
        self.assertEqual(run['stop_after'], 'scenarios')
        self.assertEqual(run['_interrupt_id'], 'scenario-confirmation')
        self.assertIn(change['id'], run['_source_ids'])
        self.assertEqual(run['_source_roles'][change['id']], 'change')
        self.assertGreater(run.get('input_version', 0), 0)
        self.assertEqual(len(self.store.runs(self.chat['id'])), 2)


if __name__ == '__main__':
    unittest.main()
