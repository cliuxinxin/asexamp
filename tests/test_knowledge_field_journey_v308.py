"""Real HTTP + native tool calling: source disclosure, scoped rebuild, field sync/export."""
import io
import time
from openpyxl import load_workbook
from test_native_journey_v300 import NativeJourney, SHARED_RULE, native_journey


def await_gate(j, run_id, kind):
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline:
        state = j.snapshot()
        run = next(row for row in state['runs'] if row['id'] == run_id)
        assert run['status'] != 'failed', run
        if run['status'] == 'waiting':
            prompt = state['conversation_prompt']
            assert prompt['kind'] == kind, prompt
            return run, j.artifact(prompt['artifact_id']), prompt
        time.sleep(.02)
    raise AssertionError(state)


def test_shared_fact_disable_native_resume_field_preview_and_real_excel(native_journey):
    origin = native_journey
    saved = origin.turn('保存这条项目澄清供同事使用：' + SHARED_RULE, 'add_knowledge_tool',
        {'content': SHARED_RULE, 'role': 'clarification', 'share': True, 'confirmed': True})
    shared_id = saved['actions'][0]['result']['source']['id']
    j = NativeJourney(origin.client, origin.app, origin.gateway)
    original_generate = j.gateway.generate_native

    async def generate(task, context, schema, instruction):
        result = await original_generate(task, context, schema, instruction)
        if task == 'generate_cases':
            result['items'][0]['test_data'] = 'valid-account'
        return result

    j.gateway.generate_native = generate
    j.turn('根据上传需求生成并评审用例，逐步确认', 'start_pipeline_tool', {'mode': 'hitp', 'stop_after': 'review'})
    old_run, old_analysis, _ = j.gate('strategy_review')
    assert SHARED_RULE in old_analysis['items'][0]['description']
    note = next(message for message in j.snapshot()['messages']
        if any(part['type'] == 'project_knowledge' for part in message.get('metadata', {}).get('turn_response', {}).get('parts', [])))
    facts = note['metadata']['turn_response']['parts'][0]['facts']
    assert facts[0]['source_id'] == shared_id and facts[0]['origin_chat_title']
    assert any(row['classification'] == 'project_knowledge' for row in old_analysis['report']['source_provenance'])
    context_path = '/api/projects/' + j.project['id'] + '/shared-context'
    context = j.client.get(context_path, params={'chat_id': j.chat['id']}).json()
    disabled = j.client.patch('/api/chats/' + j.chat['id'] + '/project-knowledge/' + shared_id,
        json={'enabled': False, 'expected_version': context['preference_version']})
    assert disabled.status_code == 200, disabled.text
    assert disabled.json()['requires_rebuild']
    prompt = j.snapshot()['conversation_prompt']
    resumed = j.turn('同意，按本次知识选择重新理解需求', 'resume_pipeline_tool',
        {'run_id': old_run['id']}, reply=prompt)
    new_run_id = resumed['actions'][0]['result']['run_id']
    assert new_run_id != old_run['id']
    _, current, prompt = await_gate(j, new_run_id, 'strategy_review')
    assert SHARED_RULE not in str(current['items'])
    assert shared_id not in str(current['report']['source_provenance'])
    assert j.artifact(old_analysis['id']) == old_analysis
    for next_kind in ('scenario_review', 'case_result_review'):
        j.turn('同意，继续', 'resume_pipeline_tool', {'run_id': new_run_id}, reply=prompt)
        _, current, prompt = await_gate(j, new_run_id, next_kind)
    j.turn('评审同意，完成', 'resume_pipeline_tool', {'run_id': new_run_id}, reply=prompt)
    drift_path = '/api/chats/' + j.chat['id'] + '/field-drift'
    detected = j.client.get(drift_path).json()
    assert 'test_data' in {row['field'] for row in detected['candidates']}
    body = {key: detected[key] for key in ('artifact_id', 'revision', 'head_revision', 'profile_id', 'profile_version')}
    staged = j.client.post(drift_path + '/propose', json={**body, 'fields': ['test_data']})
    assert staged.status_code == 200, staged.text
    change_path = '/api/chats/' + j.chat['id'] + '/profile-change'
    preview = j.client.get(change_path, params={'prompt_id': staged.json()['pending'][0]['id']}).json()
    applied = j.client.post(change_path + '/apply', json={'prompt_id': preview['prompt_id'],
        'expected_version': preview['expected_version'], 'selected_keys': ['excel_columns']})
    assert applied.status_code == 200, applied.text
    excel = j.client.get('/api/artifacts/' + current['id'] + '/export', params={'profile_id': detected['profile_id']})
    assert excel.status_code == 200, excel.text
    sheet = load_workbook(io.BytesIO(excel.content)).active
    columns = applied.json()['profile']['config']['excel_columns']
    header = next(column['header'] for column in columns if column['field'] == 'test_data')
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index(header) + 1).value == 'valid-account'
    other_context = j.client.get(context_path, params={'chat_id': origin.chat['id']}).json()
    assert next(row for row in other_context['clarifications'] if row['id'] == shared_id)['enabled_in_chat']
