"""Review tables and Excel must use the same field projection."""
import copy
import io

import pytest
from openpyxl import load_workbook

from tcg.documents import export_cases
from tcg.schemas import DomainError
from tcg.table_projection import case_table_projection


def example_artifact():
    return {
        'type': 'cases',
        '_profile': {'sheet_name': 'Review / custom', 'excel_columns': [
            {'field': 'expected', 'header': '预期结果'},
            {'field': 'test_case_description', 'header': '用例描述'},
            {'field': 'steps', 'header': '操作步骤'},
            {'field': 'payload', 'header': '输入数据', 'required': False},
            {'field': 'execution_result', 'header': '实际结果', 'value_source': 'manual'},
            {'field': 'attempts', 'header': '重试次数', 'value_source': 'default', 'default_value': 0},
            {'field': 'enabled', 'header': '启用', 'value_source': 'default', 'default_value': False},
            {'field': 'formula_text', 'header': '=公式说明', 'required': False},
        ]},
        'items': [
            {'id': 'TC-1', 'title': 'Original', 'case_description': '使用已有描述别名',
             'payload': {'用户': ['甲', 3], 'enabled': False}, 'formula_text': ' =SUM(A1:A2)\x00',
             'steps': [{'action': '打开页面\n点击登录', 'expected': '显示登录框'},
                       {'action': '提交凭据', 'expected': '显示用户首页'}]},
            {'id': 'TC-2', 'description': 'Current description', 'payload': ['one', 'two'],
             'execution_result': '人工已填写', 'attempts': 4, 'enabled': True,
             'steps': [{'action': '@输入文本', 'expected': '保留输入'}]},
        ],
    }


@pytest.mark.parametrize('layout', ['case', 'step'])
def test_projected_cells_are_exactly_the_exported_cells(layout):
    artifact = example_artifact()
    before = copy.deepcopy(artifact)
    projection = case_table_projection(artifact, layout)
    workbook = load_workbook(io.BytesIO(export_cases(artifact, layout)))
    sheet = workbook.active
    actual = [[cell.value if cell.value is not None else '' for cell in row]
              for row in sheet.iter_rows()]
    assert actual == [[column['header'] for column in projection['columns']]] + [
        row['cells'] for row in projection['rows']]
    assert projection['layout'] == layout
    assert [column['field'] for column in projection['columns']] == [
        'expected', 'description', 'steps', 'payload', 'execution_result',
        'attempts', 'enabled', 'formula_text']
    assert actual[0][-1] == "'=公式说明"
    assert actual[1][1] == '使用已有描述别名'
    assert actual[1][3] == '{"用户": ["甲", 3], "enabled": false}'
    assert actual[1][5:7] == ['0', 'False']
    assert actual[1][-1] == "' =SUM(A1:A2)"
    assert all(cell.data_type != 'f' for row in sheet.iter_rows() for cell in row)
    assert artifact == before
    assert sheet.title == 'Review _ custom'
    workbook.close()


def test_step_projection_keeps_zero_based_step_identity_and_original_selection_order():
    artifact = example_artifact()
    projection = case_table_projection(artifact, 'step', ['TC-2', 'TC-1'])
    assert [(row['item_id'], row['step_index']) for row in projection['rows']] == [
        ('TC-1', 0), ('TC-1', 1), ('TC-2', 0)]
    assert projection['rows'][1]['cells'][0] == '2. 显示用户首页'
    assert projection['rows'][1]['cells'][2] == '2. 提交凭据'
    selected = case_table_projection(artifact, 'case', ['TC-2'])
    assert [(row['item_id'], row['step_index']) for row in selected['rows']] == [('TC-2', None)]


def test_editing_projection_is_available_when_required_fields_are_missing():
    artifact = example_artifact()
    artifact['_profile']['excel_columns'].insert(0, {'field': 'new_design_field', 'header': '待补充'})
    projection = case_table_projection(artifact)
    assert projection['rows'][0]['cells'][0] == ''
    with pytest.raises(DomainError, match='必填内容'):
        export_cases(artifact)


@pytest.mark.parametrize('field', ['_private', 'refs', 'source_ids', 'report', 'run_id'])
def test_projection_does_not_expose_internal_columns(field):
    artifact = example_artifact()
    artifact['_profile']['excel_columns'] = [{'field': field, 'header': 'Forbidden'}]
    with pytest.raises(DomainError, match='列映射无效'):
        case_table_projection(artifact)


def test_defaults_and_invalid_selection_match_export_contract():
    artifact = {'type': 'cases', 'items': [{'id': 'TC-1', 'title': 'Login', 'type': 'Business',
        'priority': 'P1', 'preconditions': '', 'steps': [{'action': 'Login', 'expected': 'Success'}]}]}
    projection = case_table_projection(artifact)
    assert [column['header'] for column in projection['columns']] == [
        'Case ID', 'Title', 'Type', 'Priority', 'Preconditions', 'Steps', 'Expected Result']
    for invalid in ([], ['missing']):
        with pytest.raises(DomainError, match='ID 无效'):
            case_table_projection(artifact, selected=invalid)
    with pytest.raises(DomainError, match='layout'):
        case_table_projection(artifact, 'unknown')
