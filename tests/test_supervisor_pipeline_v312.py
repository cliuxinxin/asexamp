"""Supervisor continuations across real Pipeline gates and Profile approval HTTP."""
import asyncio
import io
import time
from collections import Counter

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tcg.main import create_app
from test_native_journey_v300 import NativeJourneyGateway
from test_supervisor_http_v312 import QueueGateway, post, seed, wait_files


def execution_plan(*steps):
    return {'title': '生成并按已确认内容导出', 'steps': [
        {'capability': capability, 'instruction': instruction}
        for capability, instruction in steps]}


def plan_state(client, chat_id, plan_id):
    response = client.get(f'/api/chats/{chat_id}/plans/{plan_id}')
    assert response.status_code == 200, response.text
    return response.json()


def await_gate(client, chat_id, kind):
    deadline = time.monotonic() + 8
    snapshot = None
    while time.monotonic() < deadline:
        snapshot = client.get('/api/chats/' + chat_id).json()
        runs = snapshot['runs']
        assert len(runs) == 1, snapshot
        assert runs[0]['status'] != 'failed', runs[0]
        prompt = snapshot.get('conversation_prompt')
        if runs[0]['status'] == 'waiting' and prompt and prompt['kind'] == kind:
            return runs[0], prompt
        time.sleep(.02)
    raise AssertionError(snapshot)


def test_plan_waits_for_each_real_human_gate_then_exports_approved_cases(tmp_path):
    gateway = QueueGateway()
    gateway.business_model = NativeJourneyGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        project = client.get('/api/projects').json()[0]
        created = client.post('/api/projects/' + project['id'] + '/chats', json={'title': '生成后导出'})
        assert created.status_code == 200, created.text
        chat = created.json()
        upload = client.post('/api/chats/' + chat['id'] + '/sources',
            files={'file': ('login.md', '注册用户输入有效账号密码后登录成功。'.encode(), 'text/markdown')},
            data={'role': 'primary'})
        assert upload.status_code == 200, upload.text
        gateway.plans = [execution_plan(
            ('pipeline_start', '根据上传需求生成用例，逐步人工确认，并包含 AI 评审。'),
            ('export', '等主流程结束，导出已经确认评审的最终用例。'))]
        gateway.calls = [('start_pipeline_tool', {'mode': 'hitp'}), ('export_artifact_tool', {})]
        initial, _ = post(client, chat, 'start-and-export', '根据需求生成用例，每一步让我确认，最后导出 Excel。',
            source_ids=[upload.json()['id']])
        assert initial['status'] == 'succeeded', initial
        part = next(p for p in initial['parts'] if p['type'] == 'execution_plan')
        stages = [('strategy_review', {'understand_requirements': 1}),
            ('scenario_review', {'understand_requirements': 1, 'generate_scenarios': 1}),
            ('case_result_review', {'understand_requirements': 1, 'generate_scenarios': 1,
                                    'generate_cases': 1, 'review_cases': 1})]
        run_id = None
        for index, (kind, counts) in enumerate(stages):
            run, prompt = await_gate(client, chat['id'], kind)
            run_id = run['id']
            # Wait through a Supervisor watcher tick: an available draft must not be exported.
            time.sleep(1.05)
            snapshot = client.get('/api/chats/' + chat['id']).json()
            assert snapshot['conversation_prompt']['id'] == prompt['id']
            assert Counter(task for task, _ in gateway.business_model.generations) == counts
            assert not app.state.store.list('frozen_export', chat_id=chat['id'])
            assert gateway.executed == ['start_pipeline_tool']
            state = plan_state(client, chat['id'], part['plan_id'])
            assert state['status'] == 'waiting_pipeline', state
            assert state['steps'][1]['status'] == 'pending', state
            if kind == 'case_result_review':
                draft = app.state.store.get('artifact', prompt['artifact_id'])
                assert not draft['items'][0]['title'].startswith('已评审：')
            result, _ = post(client, chat, 'approve-' + str(index), '同意', reply_kind='confirm',
                reply_to=prompt['id'], command={'name': 'workflow.resume',
                    'arguments': {'run_id': run_id, 'action': 'approved'}})
            assert result['status'] == 'succeeded', result
        files = wait_files(client, chat, result)
        assert len(files) == 1
        completed = app.state.store.run(run_id)
        assert completed['status'] == 'completed', completed
        artifact = app.state.store.get('artifact', completed['current_artifact_id'])
        assert artifact['type'] == 'cases'
        assert artifact['items'][0]['title'].startswith('已评审：')
        records = app.state.store.list('frozen_export', chat_id=chat['id'])
        assert len(records) == 1
        assert records[0]['artifact_id'] == artifact['id']
        assert records[0]['revision'] == artifact['revision']
        sheet = load_workbook(io.BytesIO(client.get(files[0]['url']).content)).active
        rows = list(sheet.values)
        assert len(rows) == 2 and artifact['items'][0]['title'] in rows[1], rows
        assert len(gateway.planner_inputs) == 1
        assert gateway.executed == ['start_pipeline_tool', 'export_artifact_tool']
        assert plan_state(client, chat['id'], part['plan_id'])['status'] == 'completed'


def test_profile_http_approval_continues_export_once_with_new_profile_version(tmp_path):
    gateway = QueueGateway()
    app = create_app(tmp_path, gateway)
    with TestClient(app) as client:
        chat, cases = asyncio.run(seed(app.state.store, gateway))
        profile = app.state.store.list('profile', project_id=chat['project_id'])[0]
        gateway.plans = [execution_plan(
            ('profile_edit', '给导出模板增加人工填写的执行人列，先预览确认。'),
            ('export', '确认 Profile 更改以后，按新模板导出当前用例。'))]
        gateway.calls = [('modify_profile_tool', {'profile_id': profile['id'],
            'upsert_columns': [{'field': 'tester', 'header': '执行人', 'value_source': 'manual'}]}),
            ('export_artifact_tool', {'artifact_ids': [cases['id']], 'profile_id': profile['id']})]
        initial, _ = post(client, chat, 'profile-export', '给导出模板加执行人列，然后导出这份用例。',
            artifact_id=cases['id'], artifact_revision=cases['revision'], profile_id=profile['id'])
        assert initial['status'] == 'needs_confirmation', initial
        part = next(p for p in initial['parts'] if p['type'] == 'execution_plan')
        prompt = initial['pending'][0]
        assert prompt['kind'] == 'profile'
        assert not app.state.store.list('frozen_export', chat_id=chat['id'])
        assert app.state.store.get('profile', profile['id']) == profile
        approval = {'prompt_id': prompt['id'], 'expected_version': profile['version'],
                    'selected_keys': ['excel_columns']}
        applied = client.post('/api/chats/' + chat['id'] + '/profile-change/apply', json=approval)
        assert applied.status_code == 200, applied.text
        files = wait_files(client, chat, applied.json())
        assert len(files) == 1
        records = app.state.store.list('frozen_export', chat_id=chat['id'])
        assert len(records) == 1, records
        assert records[0]['artifact_id'] == cases['id'] and records[0]['revision'] == cases['revision']
        assert records[0]['profile_id'] == profile['id']
        assert records[0]['profile_version'] == profile['version'] + 1
        assert '执行人' in next(load_workbook(io.BytesIO(client.get(files[0]['url']).content)).active.values)
        assert app.state.store.get('artifact', cases['id'])['revision'] == cases['revision']
        assert plan_state(client, chat['id'], part['plan_id'])['status'] == 'completed'
        assert len(gateway.planner_inputs) == 1
        duplicate = client.post('/api/chats/' + chat['id'] + '/profile-change/apply', json=approval)
        assert duplicate.status_code == 409, duplicate.text
        assert len(app.state.store.list('frozen_export', chat_id=chat['id'])) == 1
        assert gateway.executed == ['modify_profile_tool', 'export_artifact_tool']
